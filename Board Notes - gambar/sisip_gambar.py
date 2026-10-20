"""Sisipkan gambar ke catatan board .docx tanpa menimpa file asli (hasil: <nama>_v2.docx).

Pemakaian:
  python sisip_gambar.py msk  <Catatan_Muskuloskeletal.docx>
  python sisip_gambar.py info <file.docx> <spec.json>
spec.json (mode info): [{"section": "...", "anchor": "...", "image": "...", "caption": "..."}]
  -> gambar + caption disisipkan tepat SEBELUM paragraf `anchor` pertama yang muncul setelah
     paragraf yang diawali `section` (biasanya anchor = judul subbab berikutnya).
"""
import json, sys, pathlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = pathlib.Path(__file__).parent
NAVY = RGBColor(0x1F, 0x3A, 0x5F)


def find_anchor(doc, section, anchor):
    paras = doc.paragraphs
    start = next(i for i, p in enumerate(paras) if p.text.strip().startswith(section))
    for p in paras[start + 1:]:
        if p.text.strip().startswith(anchor):
            return p
    raise LookupError(f"anchor '{anchor}' tidak ditemukan setelah '{section}'")


def label(par, text, size=10, bold=True, italic=False, color=NAVY):
    r = par.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return r


def usable_width(doc):
    s = doc.sections[0]
    return s.page_width - s.left_margin - s.right_margin


def insert_msk(doc_path):
    sys.path.insert(0, str(HERE))
    from msk_captions import CAP, GROUPS
    src = json.loads((HERE / "msk" / "sumber_video.json").read_text())
    src["wartenberg_frame"] = src["wartenberg"]
    doc = Document(doc_path)
    for anchor, where, keys in GROUPS:
        section = where.split(":", 1)[1]
        a = find_anchor(doc, section, anchor)
        t = a.insert_paragraph_before()
        label(t, "Gambar tes provokasi", 11)
        for k in keys:
            name, target, teknik, positif, channel = CAP[k]
            tbl = doc.add_table(rows=1, cols=2)
            try:
                tbl.style = "Table Grid"
            except KeyError:
                pass
            c0, c1 = tbl.rows[0].cells
            c0.width, c1.width = Cm(7.0), Cm(9.5)
            c0.paragraphs[0].add_run().add_picture(str(HERE / "msk" / f"{k}.jpg"), width=Cm(6.6))
            p = c1.paragraphs[0]; label(p, name, 10.5)
            for lab, val in (("Target: ", target), ("Teknik: ", teknik), ("Positif bila: ", positif)):
                q = c1.add_paragraph(); label(q, lab, 9, color=None); label(q, val, 9, bold=False, color=None)
            q = c1.add_paragraph()
            label(q, f"Sumber gambar: {channel}, YouTube (youtu.be/{src[k]})", 8, bold=False, italic=True, color=RGBColor(0x55, 0x55, 0x55))
            a._p.addprevious(tbl._tbl)
            a.insert_paragraph_before().paragraph_format.space_after = Pt(2)
    return doc


def insert_info(doc_path, spec_path):
    spec = json.loads(pathlib.Path(spec_path).read_text(encoding="utf-8"))
    doc = Document(doc_path)
    w = usable_width(doc)
    for it in spec:
        a = find_anchor(doc, it["section"], it["anchor"])
        p = a.insert_paragraph_before(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(HERE / it["image"]), width=w)
        c = a.insert_paragraph_before(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label(c, it["caption"], 8.5, bold=False, italic=True, color=RGBColor(0x55, 0x55, 0x55))
    return doc


if __name__ == "__main__":
    mode, path = sys.argv[1], pathlib.Path(sys.argv[2])
    doc = insert_msk(path) if mode == "msk" else insert_info(path, sys.argv[3])
    out = path.with_name(path.stem + "_v2.docx")
    doc.save(out)
    print("saved", out)
