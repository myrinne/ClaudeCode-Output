import re, sys, difflib
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import merge

OUT = sys.argv[1]
TOPICS = [('TOX','Toksikologi Industri'),('HI','Higiene Industri, Batas Pajanan & Biomonitoring'),
          ('PARU','Penyakit Paru Akibat Kerja & Pneumokoniosis'),('KULIT','Dermatosis Akibat Kerja'),
          ('BISING','Kebisingan, Pendengaran & Getaran'),('FIS','Bahaya Fisik: Panas, Dingin, Radiasi, Tekanan'),
          ('BIO','Bahaya Biologi & Penyakit Infeksi'),('REPRO','Kesehatan Reproduksi'),
          ('ERGO','Ergonomi & Muskuloskeletal'),('PSI','Psikososial & Kesehatan Jiwa Kerja'),
          ('EPI','Epidemiologi & Statistik'),('FTW','Diagnosis PAK, Kelaikan Kerja, RTW & Kecacatan'),
          ('K3','Manajemen K3, Pengendalian & Pencegahan'),('REG','Regulasi & Jaminan Sosial'),('LING','Kesehatan Lingkungan')]
TNAME = dict(TOPICS)
GREY = RGBColor(0x59,0x59,0x59); ACC = RGBColor(0x1F,0x3A,0x5F); RED = RGBColor(0x9C,0x27,0x27)
LBL = merge.LABEL
def lab(e): return LBL.get(e, e)

items = merge.merge()
# ---------- case grouping ----------
STEMKEYS = {'pengecat mobil':'car painter','rinitis alergi pekerja kantor':'allergic rhinitis',
            'sipir penjara dengan TB':'prison','pekerja perakitan elektronik':'electronics assembly','Rudi':'Rudi works'}
stems = {k: next(x for x in items if x['type']=='stem' and v in x['Q']) for k,v in STEMKEYS.items()}
for x in items:
    m = re.match(r'^\(Kasus ([^)]*)\)\s*', x['Q'])
    x['case'] = None
    if m and m.group(1) in stems:
        x['case'] = m.group(1); x['Q'] = x['Q'][m.end():]
        x['topic'] = stems[x['case']]['topic']
        # strip repeated case narrative if the representative carried it
        st = stems[x['case']]['Q'].split(':',1)[1].strip()[:60]
        if st[:40] in x['Q']:
            q = x['Q']; i = q.rfind('. ', 0, len(q)-5)
            # keep only last sentence(s) after narrative: take text after the stem narrative's final sentence
            x['Q'] = re.split(r'(?<=\.)\s+(?=[A-Z][^.]*\?$)', q)[-1]
for x in items:
    if x['Q'].startswith('(Kasus radiografer yang sama)'): x['topic'] = 'FTW'
# variant note for crane/panel operator
for x in items:
    if 'panel operator' in x['Q'] and 'Okt 2024' in x['exams']:
        x['N'].append('Versi Okt 2024: operator crane 36 tahun dengan kelemahan tungkai kiri dan kesadaran somnolen; opsi: DHF, typhoid fever, chikungunya, heat stroke, heat cramp.')
for x in items:
    x['first'] = min(merge.ORDER.index(e) for e in x['exams'])
    x['ord'] = min(x['members'])

def order_items(lst):
    keyed = []
    for x in lst:
        if x['case']:
            st = stems[x['case']]; keyed.append(((min(merge.ORDER.index(e) for e in st['exams']), min(st['members'])), x['ord'], x))
        else:
            keyed.append(((x['first'], x['ord']), x['ord'], x))
    return [k[2] for k in sorted(keyed, key=lambda k:(k[0],k[1]))]

# ---------- docx helpers ----------
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
for side in ('left_margin','right_margin'): setattr(sec, side, Cm(2.3))
sec.top_margin = sec.bottom_margin = Cm(2.2)
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
st.paragraph_format.space_after = Pt(2); st.paragraph_format.line_spacing = 1.08
for lvl, size in ((1,16),(2,13),(3,11.5)):
    h = doc.styles[f'Heading {lvl}']; h.font.name='Calibri'; h.font.size=Pt(size); h.font.color.rgb = ACC; h.font.bold=True
    h.element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); h.element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri')
    h.paragraph_format.space_before = Pt(14 if lvl==1 else 10); h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True

def para(text='', size=None, bold=False, italic=False, color=None, indent=0, after=2, before=0, keep=False, align=None):
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text); r.bold=bold; r.italic=italic
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
    pf = p.paragraph_format; pf.left_indent = Cm(indent); pf.space_after = Pt(after); pf.space_before = Pt(before)
    if keep: pf.keep_with_next = True
    if align: p.alignment = align
    return p

def shade(p, fill='EEF2F7'):
    pPr = p._p.get_or_add_pPr(); shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),fill); pPr.append(shd)
    bdr = OxmlElement('w:pBdr')
    for s in ('left',):
        b = OxmlElement(f'w:{s}'); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'18'); b.set(qn('w:space'),'6'); b.set(qn('w:color'),'1F3A5F'); bdr.append(b)
    pPr.append(bdr)

def toc():
    p = doc.add_paragraph(); r = p.add_run()
    for t,txt in (('begin',None),(None,'TOC \\o "1-2" \\h \\z \\u'),('separate',None),(None,None),('end',None)):
        if t:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), t); r._r.append(fc)
        elif txt:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'),'preserve'); it.text = txt; r._r.append(it)
        else:
            r2 = p.add_run('Klik kanan di sini → Update Field untuk menampilkan daftar isi.'); r2.italic=True; r2.font.color.rgb=GREY
            r = p.add_run()

def page_numbers():
    p = sec.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(); r.font.size = Pt(9); r.font.color.rgb = GREY
    for t,txt in (('begin',None),(None,'PAGE'),('end',None)):
        if t: fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'),t); r._r.append(fc)
        else: it = OxmlElement('w:instrText'); it.text = txt; r._r.append(it)
    h = sec.header.paragraphs[0]; h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = h.add_run('Kompilasi Soal Ujian Board Sp.Ok — TURUNAN SENIOR'); hr.font.size=Pt(8.5); hr.font.color.rgb=GREY

def tagline(x, indent=0.6):
    p = para(indent=indent, after=1)
    r = p.add_run('Keluar di: ' + ' · '.join(lab(e) for e in x['exams'])); r.italic=True; r.font.size=Pt(8.5); r.font.color.rgb=GREY
    if len(x['exams']) > 1:
        r2 = p.add_run(f'   (keluar {len(x["exams"])}×)'); r2.bold=True; r2.font.size=Pt(8.5); r2.font.color.rgb=RED

def letter_of(x):
    if not x['O'] or not x['A']: return None
    a = re.sub(r'\s*\([^)]*\)\s*$', '', x['A'].strip()).rstrip('.').lower()
    best, bi = 0, None
    for i,o in enumerate(x['O']):
        ol = o.strip().rstrip('.').lower()
        r = 1.0 if (len(a) >= 3 and ol.startswith(a)) else difflib.SequenceMatcher(None, a, ol).ratio()
        if r > best: best, bi = r, i
    return chr(65+bi) if best >= 0.75 else None

KEY = []   # (chapter, number, x)
GREEN = RGBColor(0x1E,0x6B,0x34)
from collections import Counter
CORR = []
def load_ver():
    import os
    src = open(os.path.join(os.path.dirname(os.path.abspath(__file__)),'verifikasi.txt'), encoding='utf-8').read()
    out = {}
    for blk in re.split(r'(?m)^@ ', src)[1:]:
        lines = blk.strip().split('\n'); d = {}
        for ln in lines[1:]:
            m = re.match(r'^([EKRXS]): ?(.*)$', ln)
            if m: d[m.group(1)] = m.group(2).strip()
        out[int(lines[0].strip())] = d
    return out
VER = load_ver()
def letter_from(x, text):
    if not x['O']: return None
    a = re.sub(r'\s*\([^)]*\)\s*$', '', text.strip()).rstrip('.').lower()
    best, bi = 0, None
    for i,o in enumerate(x['O']):
        ol = o.strip().rstrip('.').lower()
        r = 1.0 if (len(a) >= 3 and (ol.startswith(a) or a.startswith(ol))) else difflib.SequenceMatcher(None, a, ol).ratio()
        if r > best: best, bi = r, i
    return chr(65+bi) if best >= 0.8 else None
def ver_status(x, v):
    if v.get('S'): return v['S']
    E = v['E']
    if E.startswith('Tidak dapat diverifikasi'): return 'tidak'
    if not x['A']: return 'baru'
    ls, lv = letter_of(x), letter_from(x, E)
    if ls and lv: return 'sesuai' if ls == lv else 'koreksi'
    return 'sesuai' if merge.same_ans(x['A'], E) else 'koreksi'
def question(n, x, chapter):
    p = para(after=1, before=5, keep=True)
    r = p.add_run(f'{n}. '); r.bold = True; r.font.color.rgb = ACC
    p.add_run(x['Q'])
    p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.first_line_indent = Cm(-0.6)
    for i,o in enumerate(x['O']):
        q = para(f'{chr(65+i)}.  {o}', indent=1.2, after=0, keep=True)
        q.paragraph_format.first_line_indent = Cm(-0.6)
    if x['type']=='mcq' and not x['O']:
        para('(opsi jawaban tidak tercatat di sumber)', size=9, italic=True, color=GREY, indent=0.6, after=0, keep=True)
    for j,s in enumerate(x['S']):
        q = para(f'{chr(97+j)}.  {s}', indent=1.2, after=0, keep=True); q.paragraph_format.first_line_indent = Cm(-0.6)
    tagline(x)
    KEY.append((chapter, n, x))

def case_box(stem, nums):
    txt = stem['Q'].split(':',1)[1].strip()
    p = para(before=6, after=3, keep=True); shade(p)
    p.paragraph_format.left_indent = Cm(0.2); p.paragraph_format.right_indent = Cm(0.2)
    r = p.add_run(f'Kasus untuk soal nomor {nums[0]}–{nums[-1]}\n' if len(nums)>1 else f'Kasus untuk soal nomor {nums[0]}\n'); r.bold=True; r.font.color.rgb=ACC; r.font.size=Pt(10)
    p.add_run(txt).font.size = Pt(10)
    p2 = para(indent=0.2, after=2, keep=True)
    r = p2.add_run('Kasus ini keluar di: ' + ' · '.join(lab(e) for e in stem['exams'])); r.italic=True; r.font.size=Pt(8.5); r.font.color.rgb=GREY

def chapter(title, intro, types, chap_id, counter_start=1):
    doc.add_heading(title, level=1)
    if intro: para(intro, italic=True, color=GREY, size=9.5, after=6)
    n = counter_start
    for code, name in TOPICS:
        lst = [x for x in items if x['topic']==code and x['type'] in types]
        if not lst: continue
        doc.add_heading(f'{name}  ({len(lst)} soal)', level=2)
        lst = order_items(lst)
        i = 0
        while i < len(lst):
            x = lst[i]
            if x['case']:
                grp = [y for y in lst[i:] if y['case']==x['case']]
                case_box(stems[x['case']], list(range(n, n+len(grp))))
                for y in grp: question(n, y, chap_id); n += 1
                i += len(grp)
            else:
                question(n, x, chap_id); n += 1; i += 1
    return n

# ---------- title & intro ----------
page_numbers()
para('KOMPILASI SOAL UJIAN BOARD', size=24, bold=True, color=ACC, before=120, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
para('Kedokteran Okupasi (Sp.Ok)', size=16, color=ACC, after=18, align=WD_ALIGN_PARAGRAPH.CENTER)
para('Pilihan Ganda · Isian Singkat & Essay · OSCE & Ujian Lisan', size=12, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
para('Soal yang pernah keluar 2016–2026 (UKOM 22), dikompilasi dari folder TURUNAN SENIOR: MCQ, Soal Ujian, Review', size=10.5, color=GREY, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
nm = sum(1 for x in items if x['type']=='mcq'); ns = sum(1 for x in items if x['type'] in ('short','essay','case')); no = sum(1 for x in items if x['type'] in ('osce','oral'))
para(f'{nm} soal pilihan ganda · {ns} soal isian/essay/kasus tulis · {no} stase OSCE/ujian lisan', size=10.5, bold=True, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
para('Disusun 15 Oktober 2026', size=10, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

doc.add_heading('Daftar Isi', level=1)
toc()
doc.add_page_break()

doc.add_heading('Cara Memakai Kompilasi Ini', level=1)
for t in [
 'Hanya soal yang pernah keluar di ujian (rekaman/recall ujian asli) yang dimasukkan. Soal simulasi buatan senior/AI (Simulasi Soal Mas Rizqi 1–200, Simulasi Ujian Tulis Feb 2026, Simulation Written Exam Feb 2026, Simulasi Soal BOARD) dan bahan review tanya-jawab (Board written exam preparation/“Board Feb 2025”, Review Bahan Ujian Tulis) tidak dimasukkan.',
 'Soal yang sama dianggap satu. Setiap soal diberi label “Keluar di:” berisi semua ujian tempat soal itu muncul; label merah “(keluar n×)” menandai soal berulang — prioritas belajar.',
 'Soal dan pilihan jawaban tampil tanpa tanda jawaban. Semua kunci ada di Lampiran A (Kunci Jawaban), dengan nomor yang sama.',
 'Kunci berasal dari sumber (kunci Google Form, tanda merah/sorotan hasil diskusi angkatan, atau jawaban recall). Bila sumber berbeda pendapat, kunci ditandai “Diperdebatkan” dan semua versi ditulis. “[Perlu dicek]” menandai kunci sumber yang bertentangan dengan rujukan standar — bukan pembahasan, hanya penanda untuk diverifikasi.',
 'Verifikasi berbasis bukti (Oktober 2026): soal yang kuncinya diperdebatkan, ditandai [Perlu dicek], atau tidak punya kunci diberi kotak hijau “✔ Jawaban terverifikasi” di Lampiran A, lengkap dengan tingkat keyakinan dan sumber (guideline/textbook, artikel, regulasi, catatan belajar Anda). Kunci sumber yang ternyata keliru ditandai “KOREKSI” dan dirangkum di Lampiran C.',
 'Soal recall yang tidak lengkap tetap dimasukkan apa adanya (opsi atau kunci yang tidak tercatat ditulis demikian), karena topiknya pernah keluar.',
 'Susunan: Bab 1 Pilihan Ganda, Bab 2 Isian Singkat/Essay/Kasus Tulis, Bab 3 OSCE & Ujian Lisan; di dalam tiap bab dikelompokkan per topik. Soal kasus berseri ditampilkan di bawah kotak kasusnya.',
 'Label ujian: “Remed Agt 2022” = Remedial Written Test Agustus 2022 (file “Board sebelumnya”); “Des 2022” = ujian tulis Desember 2022; “Awal 2022” = recall “Soal Essay 2022” dari deck latihan Juni 2022; “2021 (Des)” = Ujian Nasional 11 & 18 Desember 2021 (Google Form); “2021 (recall peserta)” = rekap peserta 2021; “UKOM 22 (2026)” = Ujian Kompetensi edisi 22. Rincian file per label ada di Lampiran B.']:
    p = para(indent=0.5, after=3); p.paragraph_format.first_line_indent = Cm(-0.4); p.add_run('•  ' + t)

# exam coverage table
doc.add_heading('Cakupan per Ujian', level=2)
from collections import Counter, defaultdict
cov = defaultdict(Counter)
for x in items:
    k = 'MCQ' if x['type']=='mcq' else ('OSCE/Lisan' if x['type'] in ('osce','oral') else ('Isian/Essay/Kasus' if x['type']!='stem' else None))
    if k:
        for e in x['exams']: cov[e][k] += 1
tb = doc.add_table(rows=1, cols=4); tb.style = 'Light Grid Accent 1'; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
for c,h in zip(tb.rows[0].cells, ['Ujian','Pilihan ganda','Isian / essay / kasus','OSCE / lisan']): c.text = h
for e in merge.ORDER:
    if e not in cov: continue
    row = tb.add_row().cells
    row[0].text = lab(e); row[1].text = str(cov[e]['MCQ'] or '–'); row[2].text = str(cov[e]['Isian/Essay/Kasus'] or '–'); row[3].text = str(cov[e]['OSCE/Lisan'] or '–')
for r in tb.rows:
    for c in r.cells:
        for p in c.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs: run.font.size = Pt(9.5)
para('Angka = jumlah soal unik yang keluar di ujian tersebut (satu soal bisa terhitung di beberapa ujian).', size=8.5, italic=True, color=GREY, before=3)

# repeated table
doc.add_heading('Soal yang Paling Sering Berulang', level=2)
rep = sorted([x for x in items if len(x['exams'])>=3 and x['type']!='stem'], key=lambda x:(-len(x['exams']), x['first']))
for x in rep:
    p = para(indent=0.5, after=1); p.paragraph_format.first_line_indent = Cm(-0.4)
    r = p.add_run(f'{len(x["exams"])}×  '); r.bold=True; r.font.color.rgb=RED
    q = x['Q'] if len(x['Q'])<130 else x['Q'][:127].rsplit(' ',1)[0]+' …'
    p.add_run(q + '  ')
    r = p.add_run('(' + ', '.join(lab(e) for e in x['exams']) + ')'); r.font.size=Pt(8.5); r.font.color.rgb=GREY
doc.add_page_break()

# ---------- chapters ----------
chapter('Bab 1. Pilihan Ganda', 'Pilih satu jawaban terbaik. Kunci di Lampiran A.1.', ('mcq',), 'A.1')
doc.add_page_break()
chapter('Bab 2. Isian Singkat, Essay & Kasus Tulis', 'Jawab singkat. Kunci di Lampiran A.2. Soal kasus tulis 2016/2020 tidak memiliki kunci di sumber.', ('short','essay','case'), 'A.2')
doc.add_page_break()

# OSCE chapter with station map
doc.add_heading('Bab 3. OSCE & Ujian Lisan', level=1)
para('Stase OSCE dan kasus ujian lisan/KOMPRE yang pernah keluar. Jawaban recall (bila ada) di Lampiran A.3.', italic=True, color=GREY, size=9.5, after=6)
doc.add_heading('Peta Stase OSCE per Periode', level=2)
para('Sumber: “Simulasi OSCE.xlsx” (rekap angkatan). [M] = stase yang selalu muncul.', size=9, italic=True, color=GREY)
grid = [('[M] Promkes','X','X','X','X','X'),('[M] Muskuloskeletal, Ergonomi, JSA','X','X','X','X','X'),('[M] ILO','X','X','X','X','X'),
        ('Psikososial','X','X','','',''),('Spirometri','X','X','X','',''),('Audiometri','X','X','X','',''),
        ('Biomonitoring','X','X','','',''),('Kecacatan','X','','','',''),('Oral test','X','X','X','X','X')]
tb = doc.add_table(rows=1, cols=6); tb.style='Light Grid Accent 1'; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
for c,h in zip(tb.rows[0].cells, ['Stase','2024-1','2024-2','2025-1','2025-2','2026-1']): c.text=h
for g in grid:
    row = tb.add_row().cells
    for c,v in zip(row,g): c.text = v
for r in tb.rows:
    for c in r.cells:
        for p in c.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs: run.font.size = Pt(9.5)
para('Catatan: tanda X disalin apa adanya dari file sumber (urutan sel kosong di file mungkin tidak sejajar kolom). Stase yang tercatat di recall: ±2020 (audiometri, spirometri, muskuloskeletal, promkes APD, JSA, ILO); Okt 2024 (muskuloskeletal, spirometri, toksikologi, JSA, promkes, ILO); Okt 2025 (promkes, biomonitoring, psikososial, spirometri, ILO, JSA-muskuloskeletal); UKOM 22 (kecacatan, promkes, muskuloskeletal, spirometri, ergonomi, ILO).', size=8.5, italic=True, color=GREY, before=3, after=6)
n = 1
for kind, title in (('osce','Stase OSCE'),('oral','Ujian Lisan & KOMPRE')):
    doc.add_heading(title, level=2)
    lst = sorted([x for x in items if x['type']==kind], key=lambda x:(x['first'], [c for c,_ in TOPICS].index(x['topic']), x['ord']))
    for x in lst:
        question(n, x, 'A.3'); n += 1
doc.add_page_break()

# ---------- answer key ----------
doc.add_heading('Lampiran A. Kunci Jawaban', level=1)
para('Kunci sesuai sumber. “Diperdebatkan” = sumber berbeda; “[Perlu dicek]” = kunci sumber bertentangan dengan rujukan standar, verifikasi sebelum dihafal.', size=9, italic=True, color=GREY, after=6)
names = {'A.1':'A.1 Pilihan Ganda','A.2':'A.2 Isian Singkat, Essay & Kasus Tulis','A.3':'A.3 OSCE & Ujian Lisan'}
for ch in ('A.1','A.2','A.3'):
    doc.add_heading(names[ch], level=2)
    for c, num, x in KEY:
        if c != ch: continue
        p = para(after=1, before=1); p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.8)
        r = p.add_run(f'{num}. '); r.bold = True; r.font.color.rgb = ACC
        L = letter_of(x)
        if L:
            r = p.add_run(f'{L}'); r.bold = True
            p.add_run(f' — {x["A"].rstrip(".")}')
        elif x['A']:
            p.add_run(x['A'])
        else:
            r = p.add_run('Kunci tidak tersedia di sumber.'); r.italic = True; r.font.color.rgb = GREY
        for d in x['D']:
            q = para(indent=0.8, after=0, size=9); rr = q.add_run('Diperdebatkan: '); rr.bold=True; rr.font.color.rgb=RED; rr.font.size=Pt(9)
            q.add_run(d).font.size = Pt(9)
        for nn in x['N']:
            q = para(indent=0.8, after=0)
            if nn.startswith('[Perlu dicek]'):
                rr = q.add_run('[Perlu dicek] '); rr.bold=True; rr.font.color.rgb=RED; rr.font.size=Pt(9)
                q.add_run(nn[len('[Perlu dicek]'):].strip()).font.size = Pt(9)
            else:
                rr = q.add_run('Catatan: '); rr.italic=True; rr.font.size=Pt(9); rr.font.color.rgb=GREY
                r2 = q.add_run(nn); r2.font.size=Pt(9); r2.font.color.rgb=GREY
        v = VER.get(min(x['members']))
        if v:
            st = ver_status(x, v); CORR.append((ch, num, x, v, st))
            q = para(indent=0.8, after=0, before=1); shade(q, 'EAF4EC')
            rr = q.add_run('✔ Jawaban terverifikasi'); rr.bold=True; rr.font.size=Pt(9); rr.font.color.rgb=GREEN
            lab_st = {'sesuai':' (sesuai kunci sumber)','koreksi':' (KOREKSI kunci sumber)','baru':' (kunci baru)','tidak':''}[st]
            rr = q.add_run(lab_st + f' — keyakinan {v["K"].lower()}: '); rr.font.size=Pt(9); rr.bold = st=='koreksi'
            if st=='koreksi': rr.font.color.rgb = RED
            L = letter_from(x, v['E'])
            q.add_run((f'{L} — ' if L else '') + v['E']).font.size = Pt(9)
            if v.get('X'):
                q2 = para(indent=0.8, after=0); shade(q2, 'EAF4EC'); r3 = q2.add_run(v['X']); r3.font.size=Pt(9)
            if v.get('R') and v['R'] != '—':
                q3 = para(indent=0.8, after=1); shade(q3, 'EAF4EC')
                r4 = q3.add_run('Sumber: '); r4.italic=True; r4.font.size=Pt(8.5); r4.font.color.rgb=GREY
                r5 = q3.add_run(v['R']); r5.font.size=Pt(8.5); r5.font.color.rgb=GREY
doc.add_page_break()

# ---------- correction summary ----------
doc.add_heading('Lampiran C. Ringkasan Verifikasi Kunci', level=1)
cnt = Counter(c[4] for c in CORR)
para(f'{len(CORR)} soal diverifikasi (diperdebatkan, [Perlu dicek], atau tanpa kunci): {cnt["koreksi"]} kunci sumber dikoreksi, {cnt["sesuai"]} kunci sumber terkonfirmasi, {cnt["baru"]} kunci baru untuk soal tanpa kunci, {cnt["tidak"]} tidak dapat diverifikasi karena redaksi/opsi tidak tercatat. Sumber verifikasi: textbook/guideline (ACGIH, NIOSH, OSHA, WHO, ILO, IARC, ICRP), artikel (PubMed/StatPearls), regulasi Indonesia, dan catatan belajar Anda. Prioritaskan hafalan pada kunci dengan keyakinan tinggi.', size=9.5, after=6)
doc.add_heading('Kunci sumber yang dikoreksi', level=2)
tb = doc.add_table(rows=1, cols=5); tb.style='Light Grid Accent 1'
for c,h in zip(tb.rows[0].cells, ['No.','Soal (ringkas)','Kunci sumber','Jawaban terverifikasi','Keyakinan']): c.text = h
for ch, num, x, v, st in CORR:
    if st != 'koreksi': continue
    row = tb.add_row().cells
    q = x['Q'] if len(x['Q']) < 110 else x['Q'][:107].rsplit(' ',1)[0] + ' …'
    row[0].text = f'{ch}-{num}'; row[1].text = q; row[2].text = (x['A'] or '–')[:90]; row[3].text = v['E'][:140]; row[4].text = v['K']
for r in tb.rows:
    for c in r.cells:
        for p in c.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs: run.font.size = Pt(8.5)
para('Nomor = nomor di Lampiran A (A.1 pilihan ganda, A.2 isian, A.3 OSCE/lisan).', size=8.5, italic=True, color=GREY, before=3)
doc.add_page_break()

# ---------- sources ----------
doc.add_heading('Lampiran B. Sumber per Label Ujian', level=1)
SRC = [
 ('2016','Soal Ujian Tulis 2016.docx (Soal Ujian/-2024; salinan di MCQ/2026 Februari)'),
 ('2016-Remed','Jawaban Soal remedial ujian tulis nasional 2016.pdf (beberapa salinan identik)'),
 ('2020','Ujian Tulis Board 2020.docx; kasus tulis 2020 = kasus 2016 (juga dicatat di “UJIAN NASIONAL – Ala SANSAN.pdf”)'),
 ('±2020 (catatan peserta)','UJIAN NASIONAL – Ala SANSAN.pdf (format & stase OSCE); deck “Ujian Lisan – Review for Board 2024” (contoh kasus lisan 2020)'),
 ('2021 (Des)','UJIAN NASIONAL KEDOKTERAN OKUPASI 2021 & REMEDIAL (Google Form, 6 salinan; kunci dari “Correct answer”); Ujian Tulis Board 2021.docx; gambar “soal MCQ 2022”/image1x*.png (tangkapan layar form yang sama)'),
 ('2021 (recall peserta)','UJIAN TULIS BOARD 2021. rev.docx'),
 ('Awal 2022','Deck “latihan soal audio-spiro-essay kelp sugiharto.pptx” (Juni 2022), bagian “Soal Essay 2022”'),
 ('2022 / 2023 (lisan)','Deck “Ujian Lisan – Review for Board 2024.pptx” (SOAL BOARD 2023; contoh kasus 2022)'),
 ('Remed Agt 2022','Board sebelumnya.docx (dibuat 23 Nov 2022), dikonfirmasi foto WhatsApp 23 Agt 2022 “Remedial Written Test”; kunci dari 260129 Pagi Board sebelumnya.docx (diskusi Jan 2026); salinan catatan: Materi Board mas Okky.docx'),
 ('Des 2022','Ujian tulis MCQ board.docx/.pdf (40 MCQ; kunci dari 260126 Pagi Diskusi …); Board Desember 2022.docx & Ujian Tulis Board_Desember 2022.docx (recall + isian); Belajar bersama – Ujian Tulis Board Desember 2022.docx (kunci diskusi 2022)'),
 ('Juli 2023','soal Juli 2023.pdf (tangkapan layar Google Form beranotasi; kunci = sorotan anotator); SOAL-1.pdf (versi tanpa anotasi); review osce board Juli 2023.pdf'),
 ('Feb 2024','Soal Ujian Tulis Board 29 Feb 2024.docx (2 versi; kunci diskusi Jan 2026); soal kompilasi 2016-2024 (M. Rizqi).pdf'),
 ('Okt 2024','Board Okt 2024.pdf (MCQ, isian, stase OSCE & lisan); Rangkuman Syair Ujian Lisan & OSCE.pdf'),
 ('Okt 2024-Remed','260127 Sore UJIAN REMED Okt 2024.docx (kunci merah); Board Remed Okt 2024.docx'),
 ('Okt 2025','260126 Sore SOAL BOARD OKT 2025.docx (recall + pembahasan; beberapa opsi direkonstruksi dari pembahasan)'),
 ('Feb 2026','Board Feb 2026.docx (recall, kebanyakan tanpa opsi lengkap)'),
 ('UKOM 22 (2026)','Operan Ujian UKOM 22 2026.docx (CBT); Operan_Persiapan_Board_SpOk_UKOM22.pptx (kunci CBT, 6 stase OSCE, KOMPRE)')]
for k, v in SRC:
    p = para(indent=0.5, after=2); p.paragraph_format.first_line_indent = Cm(-0.4)
    r = p.add_run(k + ': '); r.bold = True
    p.add_run(v)
doc.add_heading('Tidak dimasukkan', level=2)
for t in ['Simulasi (bukan soal asli): 260129 Sore & 260130 Pagi Board Feb 2026 – Simulasi Soal Mas Rizqi 1–200; 260130 Malam SIMULASI UJIAN TULIS FEB 2026; 260131 Malam SIMULATION WRITTEN EXAM Feb 2026; Simulasi Soal BOARD.docx.',
          'Bahan review berbentuk tanya-jawab: Board written exam 1 (.pdf/.docx) = “Board Feb 2025.docx” = “Board Written Exam 1 Feb 2025.docx”; REVIEW BAHAN UJIAN TULIS.revisi.pdf.',
          'Ringkasan/rujukan tanpa soal: Rangkuman Kasus.pdf, OCCUPATIONAL MEDICINE BOARD GUIDELINE, PANDUAN Ujian Kompetensi Nasional, Occupational and Environmental Review (buku hasil scan), RINGKASAN UJIAN BOARD, outline anamnesis/ujian kasus, INDONESIA RAYA (skrip), COACHING UJIAN LISAN 2022 (format ujian).',
          'Rekaman video (MCQ/2026 Juli/*.mp4, Review/2025 Februari/Belajar bersama H-2.mp4) tidak ditranskripsikan.']:
    p = para(indent=0.5, after=2); p.paragraph_format.first_line_indent = Cm(-0.4); p.add_run('•  ' + t)

st_el = doc.settings.element
uf = OxmlElement('w:updateFields'); uf.set(qn('w:val'),'true'); st_el.append(uf)
doc.save(OUT)
print('saved', OUT, len(KEY), 'numbered items')
