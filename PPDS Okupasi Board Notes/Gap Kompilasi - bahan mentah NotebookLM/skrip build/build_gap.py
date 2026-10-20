"""Build 'Catatan Tambahan - Gap Kompilasi' docx (all parts in parts.py) from short/*.md.
Usage: python build_gap.py OUT.docx
Reuses mdconv/extras from the toxicology build folder. md extras: lines starting
'!PEARL ' (yellow callout), '!SUMBER ' (red callout), '!IMG file | caption' (figure),
and [[wikilinks]] (shown as plain text in docx)."""
import sys, os, re
HERE = os.path.dirname(os.path.abspath(__file__))
RAW = os.path.dirname(HERE)                                  # .../Gap Kompilasi - bahan mentah NotebookLM
BOARD = os.path.dirname(RAW)                                 # .../PPDS Okupasi Board Notes
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(BOARD, "Toksikologi - bahan mentah NotebookLM", "skrip build"))
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import mdconv as M
from extras import callout
from parts import PARTS

SHORT = os.path.join(RAW, "short")
IMGDIR = os.path.join(BOARD, "Gap Kompilasi - gambar")
OUT = sys.argv[1]


def read(name):
    with open(os.path.join(SHORT, name + ".md"), encoding="utf-8") as f:
        return f.read().replace("\r", "").split("\n")


def render(doc, lines):
    buf = []

    def flush():
        if buf:
            M.render(doc, buf, h_off=1)
            buf.clear()

    for l in lines:
        l = re.sub(r"\[\[([^\]]+)\]\]", r"catatan '\1'", l)   # wikilink -> plain text in docx
        if l.startswith("!PEARL "):
            flush()
            callout(doc, "Clinical pearl: " + re.sub(r"\*+", "", l[7:]), fill="FFF2CC")
        elif l.startswith("!IMG "):
            flush()
            name, cap = [x.strip() for x in l[5:].split("|", 1)]
            path = os.path.join(IMGDIR, name)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(path, width=Cm(9.0 if ("lux" in name or "tall" in name) else 13.0))
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cr = cp.add_run(cap)
            cr.italic = True
            cr.font.size = Pt(8.5)
        elif l.startswith("!SUMBER "):
            flush()
            callout(doc, "CATATAN SUMBER: " + re.sub(r"\*+", "", l[8:]), fill="FDE0DC")
        else:
            buf.append(l)
    flush()


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.6), Cm(27.9)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2.0)
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.font.size = Pt(10.5)
for nm, sz in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
    h = doc.styles[nm]
    h.font.name = "Calibri"
    h.font.size = Pt(sz)

fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for t, instr in (("begin", None), (None, "PAGE"), ("end", None)):
    r = fp.add_run()
    if t:
        c = OxmlElement("w:fldChar"); c.set(qn("w:fldCharType"), t); r._r.append(c)
    else:
        i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = instr; r._r.append(i)
    r.font.size = Pt(8)

doc.add_paragraph("CATATAN TAMBAHAN BOARD: GAP DARI KOMPILASI SOAL 2016 SAMPAI 2026", style="Title")
letters = ", ".join(p["letter"] for p in PARTS)
p = doc.add_paragraph(); r = p.add_run("Bagian " + letters + ": " + "; ".join(p_["letter"] + " " + p_["name"].lower() for p_ in PARTS)); r.font.size = Pt(12)
p = doc.add_paragraph()
r = p.add_run("Catatan persiapan ujian board Kedokteran Okupasi (Sp.Ok). Disusun mengikuti 'Outline Catatan Tambahan - Gap Kompilasi' dari jawaban NotebookLM (notebook Board Notes - Gap) "
              "dan, bila notebook tidak dapat membaca (salinan hasil pindai), dari teks resmi regulasi yang dibaca langsung (Permenaker No. 5 Tahun 2018 beserta lampirannya). "
              "Standar Indonesia diutamakan; sumber internasional dipakai sebagai pelengkap dan ditandai. "
              "Yang sudah ada di catatan lama (misalnya ROSA, RULA, REBA di catatan JSA; nilai ambang, BEI, IARC di Catatan Toksikologi) tidak diulang dan hanya ditunjuk. "
              "Kotak kuning = clinical pearl; kotak merah muda = catatan sumber yang perlu diverifikasi. "
              "STATUS: bagian lain menyusul sesuai outline.")
r.font.size = Pt(9.5); r.italic = True

doc.add_paragraph().add_run("Daftar isi").bold = True
for part in PARTS:
    pp = doc.add_paragraph()
    rr = pp.add_run(part["letter"] + ". " + part["name"].upper())
    rr.bold = True
    pp.paragraph_format.space_after = Pt(0)
    for title, f in part["sections"]:
        doc.add_paragraph(title, style="List Bullet").paragraph_format.space_after = Pt(0)
doc.add_page_break()

for pi, part in enumerate(PARTS):
    if pi:
        doc.add_page_break()
    doc.add_heading(part["letter"] + ". " + part["name"].upper(), level=1)
    for title, f in part["sections"]:
        doc.add_heading(title, level=2)
        render(doc, read(f))

doc.save(OUT)
print("saved", OUT)
