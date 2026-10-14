import sys, os, re
sys.path.insert(0, os.path.dirname(__file__))
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import mdconv as M
import extras_short as X

TMP = os.path.join(os.environ["CLAUDE_JOB_DIR"], "tmp")
SHORT = os.path.join(TMP, "short")
IMG = os.path.join(TMP, "infografik")
OUT = sys.argv[1]


from PIL import Image
SMALL = os.path.join(TMP, "infografik_small")
os.makedirs(SMALL, exist_ok=True)
for fn in os.listdir(IMG):
    if fn.lower().endswith(".png"):
        im = Image.open(os.path.join(IMG, fn)).convert("RGB")
        im.thumbnail((2000, 2000))
        im.save(os.path.join(SMALL, fn[:-4] + ".jpg"), quality=88, optimize=True)
IMG = SMALL


def read(name):
    with open(os.path.join(SHORT, name + ".md"), encoding="utf-8") as f:
        return f.read().replace("\r", "").split("\n")


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.6), Cm(27.9)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2.0)
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.font.size = Pt(10.5)
for nm, sz in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
    h = doc.styles[nm]
    h.font.name = "Calibri"
    h.font.size = Pt(sz)

fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for t, instr in (("begin", None), (None, "PAGE"), ("end", None)):
    r = fp.add_run()
    if t:
        c = OxmlElement("w:fldChar"); c.set(qn("w:fldCharType"), t); r._r.append(c)
    else:
        i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = instr; r._r.append(i)
    r.font.size = Pt(8)

doc.add_paragraph("CATATAN BELAJAR: TOKSIKOLOGI & BIOMONITORING", style="Title")
p = doc.add_paragraph(); r = p.add_run("Bagian 1 sampai 3A: prinsip dasar, regulasi Indonesia, dan logam berat"); r.font.size = Pt(13)
p = doc.add_paragraph(); r = p.add_run("Catatan persiapan ujian board Kedokteran Okupasi (Sp.Ok), versi ringkas. Disusun dari jawaban NotebookLM (notebook \"Toxicology - Heavy metals\") mengikuti outline. Topik 3B sampai 3F menyusul dari notebook lain."); r.font.size = Pt(9.5); r.italic = True
p = doc.add_paragraph(); r = p.add_run("Daftar isi"); r.bold = True; r.font.size = Pt(12)
toc_anchor = doc.add_paragraph()
doc.add_page_break()

doc.add_heading("1. BAGIAN AWAL: PRINSIP DASAR", level=1)
P1 = [
    ("1.1 Dosis-respons", "q01"), ("1.2 LD50 dan LC50", "q02"), ("1.3 NAB/TLV", "q03"),
    ("1.4 TWA dan penyesuaian shift panjang", "q04"), ("1.5 Campuran: rumus unity (aditif)", "q05"),
    ("1.6 Biomarker", "q06"), ("1.7 BEI", "q07"), ("1.8 Koreksi urin (kreatinin dan berat jenis)", "q08"),
    ("1.9 Biotransformasi", "q09"), ("1.10 Toksikokinetik umum", "q10"), ("1.11 Mekanisme toksisitas", "q11"),
    ("1.12 Klasifikasi karsinogen", "q12"),
]
for title, f in P1:
    doc.add_heading(title, level=2)
    M.render(doc, read(f), h_off=1)

doc.add_page_break()
doc.add_heading("2. REGULASI INDONESIA", level=1)
P2 = [
    ("2.1 Permenaker 5/2018 (NAB, IPB) dan Kepmenaker 187/1999", "q13"),
    ("2.2 PAK: Perpres 7/2019 dan Permenkes 11/2022", "q14"),
    ("2.3 Status PP 88/2019 dan PP 28/2024", "q15"),
    ("2.4 ARKL: Pedoman Ditjen PP&PL Kemenkes 2012", "q16"),
]
for title, f in P2:
    doc.add_heading(title, level=2)
    M.render(doc, read(f), h_off=1)
X.regulasi_lain(doc)

doc.add_page_break()
doc.add_heading("3. BAGIAN TENGAH: TEMPLATE PER ZAT (3A. LOGAM BERAT)", level=1)
X.template_intro(doc)
METALS = [
    ("3A.1 Kromium (Cr VI vs Cr III)", "m01_Cr", "Cr"),
    ("3A.2 Merkuri (Hg)", "m02_Hg", "Hg"),
    ("3A.3 Arsen (As)", "m03_As", "As"),
    ("3A.4 Timbal (Pb)", "m04_Pb", "Pb"),
    ("3A.5 Kadmium (Cd)", "m05_Cd", "Cd"),
    ("3A.6 Mangan (Mn)", "m06_Mn", "Mn"),
    ("3A.7 Nikel (Ni)", "m07_Ni", "Ni"),
    ("3A.8 Berilium (Be)", "m08_Be", "Be"),
    ("3A.9 Kobalt dan hard metal", "m09_Co", "Co"),
    ("3A.10 Seng / zinc (metal fume fever)", "m10_Zn", "Zn"),
    ("3A.11 Talium (Tl)", "m11_Tl", "Tl"),
]
for title, f, key in METALS:
    doc.add_page_break()
    doc.add_heading(title, level=2)
    M.render(doc, read(f), h_off=1)
    for ext in ("png", "jpg"):
        ip = os.path.join(IMG, key + "." + ext)
        if os.path.exists(ip):
            cp = doc.add_paragraph()
            cp.paragraph_format.keep_with_next = True
            cr = cp.add_run("Infografik: " + title.split(" ", 1)[1])
            cr.bold = True
            cr.font.size = Pt(9.5)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(ip, width=Cm(17.2))
            break

doc.add_page_break()
X.formula(doc)
doc.add_page_break()
X.quick_table(doc)
doc.add_page_break()
X.penutup(doc)
X.quality(doc)

heads = [p for p in doc.paragraphs if p.style.name in ("Heading 1", "Heading 2")]
bm_id = 100
entries = []
for n, h in enumerate(heads):
    name = "toc_%03d" % n
    s = OxmlElement("w:bookmarkStart"); s.set(qn("w:id"), str(bm_id + n)); s.set(qn("w:name"), name)
    e = OxmlElement("w:bookmarkEnd"); e.set(qn("w:id"), str(bm_id + n))
    h._p.insert(1 if h._p.pPr is not None else 0, s)
    h._p.append(e)
    entries.append((name, h.text, h.style.name))
prev = toc_anchor._p
for name, text, sty in entries:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Cm(0 if sty == "Heading 1" else 0.8)
    hl = OxmlElement("w:hyperlink"); hl.set(qn("w:anchor"), name); hl.set(qn("w:history"), "1")
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    if sty == "Heading 1":
        rpr.append(OxmlElement("w:b"))
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "20" if sty == "Heading 1" else "19"); rpr.append(sz)
    col = OxmlElement("w:color"); col.set(qn("w:val"), "1F3864"); rpr.append(col)
    r.append(rpr)
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text; r.append(t)
    hl.append(r)
    p._p.append(hl)
    prev.addnext(p._p)
    prev = p._p

doc.save(OUT)
print("saved", OUT)
