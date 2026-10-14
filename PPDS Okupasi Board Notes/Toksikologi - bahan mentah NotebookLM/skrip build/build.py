import sys, os, re
sys.path.insert(0, os.path.dirname(__file__))
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import mdconv as M

ANS = os.path.join(os.environ["CLAUDE_JOB_DIR"], "tmp", "answers")
OUT = sys.argv[1]


def read(name):
    with open(os.path.join(ANS, name + ".txt"), encoding="utf-8") as f:
        return f.read()


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.6), Cm(27.9)
sec.left_margin = sec.right_margin = Cm(2.4)
sec.top_margin = sec.bottom_margin = Cm(2.2)
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.font.size = Pt(10.5)
for nm, sz in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
    h = doc.styles[nm]
    h.font.name = "Calibri"
    h.font.size = Pt(sz)

# footer page number
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for t, instr in (("begin", None), (None, "PAGE"), ("end", None)):
    r = fp.add_run()
    if t:
        c = OxmlElement("w:fldChar"); c.set(qn("w:fldCharType"), t); r._r.append(c)
    else:
        i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = instr; r._r.append(i)
    r.font.size = Pt(8)

doc.add_paragraph("CATATAN BELAJAR: TOKSIKOLOGI & BIOMONITORING", style="Title")
p = doc.add_paragraph(); r = p.add_run("Bagian 1 sampai 3A: Prinsip dasar, regulasi Indonesia, dan logam berat"); r.font.size = Pt(13)
p = doc.add_paragraph(); r = p.add_run("Catatan persiapan ujian board Kedokteran Okupasi (Sp.Ok). Disusun dari jawaban NotebookLM (notebook \"Toxicology - Heavy metals\", 56 sumber) mengikuti outline yang diberikan. Topik 3B sampai 3F, formula tambahan, dan quick table zat lain menyusul dari notebook lain."); r.font.size = Pt(9.5); r.italic = True
doc.add_paragraph()
p = doc.add_paragraph(); r = p.add_run("Daftar isi"); r.bold = True; r.font.size = Pt(12)
toc = doc.add_paragraph()
for t, instr in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None), ("end", None)):
    r = toc.add_run()
    if t:
        c = OxmlElement("w:fldChar"); c.set(qn("w:fldCharType"), t)
        if t == "begin": c.set(qn("w:dirty"), "true")
        r._r.append(c)
    elif instr:
        i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = instr; r._r.append(i)
doc.add_paragraph("(Klik kanan daftar isi lalu pilih Update Field bila halaman belum terisi.)").runs[0].italic = True
doc.add_page_break()

import extras  # defined separately: formula sheet, quick tables, OSCE, vignettes, quality notes

# ---------- BAGIAN 1
doc.add_heading("1. BAGIAN AWAL: PRINSIP DASAR", level=1)
P1 = [
    ("1.1 Dosis-respons", "q01"), ("1.2 LD50 dan LC50", "q02"), ("1.3 NAB/TLV", "q03"),
    ("1.4 TWA dan penyesuaian shift panjang", "q04"), ("1.5 Campuran: rumus unity (aditif)", "q05"),
    ("1.6 Biomarker", "q06"), ("1.7 BEI", "q07"), ("1.8 Koreksi urin (kreatinin dan berat jenis)", "q08"),
    ("1.9 Biotransformasi", "q09"), ("1.10 Toksikokinetik umum", "q10"), ("1.11 Mekanisme toksisitas", "q11"),
    ("1.12 Klasifikasi karsinogen", "q12"),
]
for title, f in P1:
    doc.add_heading(title, level=2)
    body, refs = M.clean_answer(read(f))
    M.render(doc, body, h_off=1)
    M.render_refs(doc, refs)

# ---------- BAGIAN 2
doc.add_page_break()
doc.add_heading("2. REGULASI INDONESIA", level=1)
P2 = [
    ("2.1 Permenaker No. 5 Tahun 2018 (K3 lingkungan kerja, NAB, IPB) dan Kepmenaker 187/1999", "q13"),
    ("2.2 PAK: Perpres 7/2019 dan Permenkes 11/2022", "q14"),
    ("2.3 Status PP 88/2019 dan PP 28/2024", "q15"),
    ("2.4 ARKL: Pedoman Ditjen PP&PL Kemenkes 2012", "q16"),
]
for title, f in P2:
    doc.add_heading(title, level=2)
    body, refs = M.clean_answer(read(f))
    M.render(doc, body, h_off=1)
    M.render_refs(doc, refs)
extras.regulasi_lain(doc)

# ---------- BAGIAN 3
doc.add_page_break()
doc.add_heading("3. BAGIAN TENGAH: TEMPLATE PER ZAT (3A. LOGAM BERAT)", level=1)
extras.template_intro(doc)
METALS = [
    ("3A.1 Kromium (Cr VI vs Cr III)", ["q17_Cr_A", "q18_Cr_B"]),
    ("3A.2 Merkuri (Hg)", ["q19_Hg_A", "q20_Hg_B"]),
    ("3A.3 Arsen (As)", ["q21_As_A", "q22_As_B"]),
    ("3A.4 Timbal (Pb)", ["q23_Pb_A", "q24_Pb_B"]),
    ("3A.5 Kadmium (Cd)", ["q25_Cd_A", "q26_Cd_B"]),
    ("3A.6 Mangan (Mn)", ["q27_Mn_A", "q28_Mn_B"]),
    ("3A.7 Nikel (Ni)", ["q29_Ni_A", "q30_Ni_B"]),
    ("3A.8 Berilium (Be)", ["q31_Be_A", "q32_Be_B"]),
    ("3A.9 Kobalt dan hard metal", ["q33_Co"]),
]
for title, files in METALS:
    doc.add_page_break()
    doc.add_heading(title, level=2)
    for fi, f in enumerate(files):
        body, refs = M.clean_answer(read(f))
        M.render(doc, body, h_off=1)
        M.render_refs(doc, refs)
# Zn / Tl (shortened)
zt = read("q34_Zn_Tl")
idx = [m.start() for m in re.finditer(r"(?m)^# BAGIAN B", zt)]
zn_md, tl_md = (zt[:idx[0]], zt[idx[0]:]) if idx else (zt, "")
for title, md in (("3A.10 Seng / zinc (metal fume fever)", zn_md), ("3A.11 Talium (Tl)", tl_md)):
    if not md.strip():
        continue
    doc.add_page_break()
    doc.add_heading(title, level=2)
    body, refs = M.clean_answer(md if md.lstrip().startswith("---") or True else md)
    M.render(doc, body, h_off=1, skip_h_prefix=("BAGIAN",))
    M.render_refs(doc, refs)

# ---------- BAGIAN 4, 5, 6
doc.add_page_break()
extras.formula(doc)
doc.add_page_break()
extras.quick_table(doc)
doc.add_page_break()
extras.penutup(doc)
extras.quality(doc)
doc.save(OUT)
print("saved", OUT)
