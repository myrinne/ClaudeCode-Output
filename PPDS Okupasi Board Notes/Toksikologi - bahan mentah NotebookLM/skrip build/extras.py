from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import mdconv as M


def para(doc, text, size=None, bold=False, italic=False, after=4, color=None, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    M.add_inline(p, text, size=size, base_bold=bold, base_italic=italic, color=color)
    return p


def bullets(doc, items, size=None):
    for t in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        M.add_inline(p, t, size=size)


def callout(doc, text, fill="FFF2CC"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    c.text = ""
    M.shade(c, fill)
    M.add_inline(c.paragraphs[0], text, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def regulasi_lain(doc):
    doc.add_heading("2.5 Regulasi lain (belum tercakup di notebook ini)", level=2)
    callout(doc,
            "**Perhatian status regulasi.** Outline menyebut PP 88/2019 tentang Kesehatan Kerja. Menurut pembacaan NotebookLM atas PP 28/2024 "
            "(Pasal 1170), PP 88/2019 **dicabut dan dinyatakan tidak berlaku**; peraturan pelaksananya, termasuk Permenkes 11/2022, tetap berlaku "
            "sepanjang tidak bertentangan dengan PP 28/2024 (Pasal 1169 huruf w). Pengaturan kesehatan kerja kini berada di PP 28/2024 "
            "(Upaya Kesehatan Kerja, Pasal 232 sampai 244). Dasar: bagian 2.3. Mohon cek ulang pasal-pasalnya langsung pada PDF PP 28/2024 sebelum ujian.")
    bullets(doc, [
        "**Permenkes 7/2019 (kesehatan lingkungan rumah sakit)** dan **regulasi BPJS Ketenagakerjaan untuk PAK**: tidak ada di notebook toksikologi logam berat ini. Perlu diisi dari sumber lain atau notebook lain.",
        "**Pegangan cepat:** NAB/lingkungan kerja = Permenaker 5/2018 (mencabut Permenakertrans 13/2011); daftar PAK dan hak JKK sampai 3 tahun setelah hubungan kerja berakhir = Perpres 7/2019; pelayanan PAK dan 7 langkah diagnosis = Permenkes 11/2022; ARKL = Pedoman Ditjen PP&PL Kemenkes 2012.",
    ])


def template_intro(doc):
    para(doc, "Setiap zat disusun mengikuti 15 poin template. Poin 1 sampai 8 ada di **Bagian A** tiap zat; poin 9 sampai 13 plus clinical pearl ada di **Bagian B**. Poin 15 (tabel singkat per poin) disajikan sebagai Quick Table di bagian 5.", size=10)
    M.add_table(doc, [
        ["No", "Poin template", "Letak"],
        ["1", "Sifat fisik dan kimia / klasifikasi", "Bagian A"],
        ["2", "Sumber (penggunaan industri)", "Bagian A"],
        ["3", "Toksikokinetik (absorpsi, distribusi, metabolisme, ekskresi)", "Bagian A"],
        ["4", "Toksikodinamik", "Bagian A"],
        ["5", "Biomonitoring: biomarker darah/urin dan waktu sampling", "Bagian A"],
        ["6", "Waktu paruh (T1/2)", "Bagian A"],
        ["7", "Efek biologis akut", "Bagian A"],
        ["8", "Efek biologis kronis", "Bagian A"],
        ["9", "Standar exposure limit (TLV, NAB)", "Bagian B"],
        ["10", "Standar BEI / IPB", "Bagian B"],
        ["11", "Tatalaksana akut/keracunan (termasuk antidotum)", "Bagian B"],
        ["12", "Tatalaksana individu dan komunitas (7 langkah PAK, kelaikan kerja, rekomendasi manajemen)", "Bagian B"],
        ["13", "Keterangan lain (karsinogenitas, NAB udara, dst.)", "Bagian B"],
        ["14", "Clinical pearl dan jebakan ujian", "Bagian B (akhir)"],
        ["15", "Tabel singkat per poin", "Bagian 5 (Quick Table)"],
    ], font=9, widths=[1.2, 11.5, 4.2])
    callout(doc,
            "**Jebakan lintas-zat yang perlu diingat.** Angka NAB dan IPB Permenaker 5/2018 sering lebih tua daripada TLV/BEI ACGIH edisi terbaru "
            "(contoh: Cr VI larut air, NAB Permenaker 0,05 mg/m3 vs TLV-TWA ACGIH 0,0002 mg/m3). Saat ujian, sebutkan keduanya dan nyatakan mana yang menjadi acuan yuridis di Indonesia.",
            fill="FDE9D9")


def formula(doc):
    doc.add_heading("4. FORMULA WAJIB HITUNG", level=1)
    para(doc, "Ringkasan rumus. Kolom sumber menunjukkan bagian catatan yang memuat penurunan dan contoh hitungnya. Rumus yang bukan berasal dari NotebookLM diberi tanda (baku).", size=9.5, italic=True)
    M.add_table(doc, [
        ["Kelompok", "Rumus", "Keterangan / sumber"],
        ["Paparan TWA 8 jam", "TWA = Σ(Ci × Ti) / ΣTi", "Contoh: 35 ppm × 2 j, 15 × 3, 25 × 1, 5 × 2 → 150/8 = 18,75 ppm. Lihat 1.4."],
        ["Adjusted OEL shift panjang (Brief & Scala)", "RFd = (8/hd) × ((24 − hd)/16); RFw = (40/hw) × ((168 − hw)/128); OELadj = OEL8 jam × RF (pakai RF terkecil)", "Shift 12 jam: RFd = (8/12) × (12/16) = 0,50 → NAB 100 ppm menjadi 50 ppm. Lihat 1.4 (model OSHA lebih longgar, tanpa penalti jam pemulihan)."],
        ["Campuran (unity)", "Σ(Ci / TLVi) ≤ 1 untuk zat dengan organ target/efek sama", "Lihat 1.5. Efek independen dinilai terpisah, bukan dijumlah."],
        ["Non-kanker (ARKL)", "Intake = (C × R × tE × fE × Dt) / (Wb × tavg); tavg = Dt × 365 hari; RQ = Intake / RfD; inhalasi RQ = C / RfC; HI = ΣRQ", "Notasi Pedoman ARKL 2012. Setara CDI = (C × IR × EF × ED)/(BW × AT) dan HQ. RQ > 1 = tidak aman. Lihat 2.4."],
        ["Konsentrasi/dosis aman (ARKL)", "C aman = (RfD × Wb × tavg) / (R × tE × fE × Dt)", "(baku) turunan aljabar dari rumus intake."],
        ["MOS / MOE", "MOE = NOAEL / pajanan (atau BMDL / pajanan)", "(baku) tidak dibahas rinci di notebook; makin besar makin aman."],
        ["Kanker", "Intake karsinogenik: tavg = 70 tahun × 365 = 25.550 hari; ECR = Intake × CSF; atau ECR = C × IUR", "Interpretasi: Pedoman ARKL memakai batas 1E-4; rentang kerja 1E-6 sampai 1E-4. LADD/LADI = intake yang dirata-ratakan seumur hidup (baku). Lihat 2.4."],
        ["Biomonitoring: koreksi kreatinin", "µg/g kreatinin = konsentrasi urin (µg/L) / kreatinin (g/L)", "Diterima bila kreatinin 0,3 sampai 3 g/L. Lihat 1.8."],
        ["Biomonitoring: koreksi berat jenis", "C koreksi = C ukur × (SG ref − 1,000) / (SG sampel − 1,000)", "SG ref lazim 1,024 (kadang 1,020); sampel layak bila 1,010 sampai 1,030. Contoh 10 µg/L pada SG 1,012 → 20 µg/L. Lihat 1.8."],
        ["Steady state", "Cukup ≈ 5 × T1/2 (≈ 97% dari Css)", "Lihat 1.10."],
        ["Kinetik: T1/2, sisa, akumulasi", "T1/2 = 0,693 / Kel = 0,693 × Vd / Cl; fraksi sisa setelah n × T1/2 = (1/2)^n; faktor akumulasi = 1 / (1 − e^(−Kel × τ))", "Lihat 1.10. Faktor akumulasi dan fraksi sisa bersifat (baku)."],
    ], font=8.5, widths=[3.3, 7.0, 6.6])


def quick_table(doc):
    doc.add_heading("5. QUICK TABLE AKHIR (3A. LOGAM BERAT)", level=1)
    para(doc, "Kolom: Zat | Gejala klinis | Biomarker spesifik (sampel + waktu) | Antidotum / terapi kunci. Angka diambil dari jawaban NotebookLM di bagian 3A; periksa catatan kualitas di akhir dokumen untuk nilai yang berbeda antar sumber.", size=9.5, italic=True)
    rows = [
        ["Zat", "Gejala klinis", "Biomarker spesifik (sampel + waktu)", "Antidotum / terapi kunci"],
        ["**Kromium (Cr VI)**", "Ulkus/perforasi septum nasi, chrome holes, dermatitis kontak alergi, asma, kanker paru dan sinonasal", "Kromium total urin, akhir shift akhir minggu kerja. IPB Permenaker/ACGIH lama 25 µg/L (peningkatan selama shift 10 µg/L); ACGIH 2026 berbasis populasi 0,7 µg/L. Cr eritrosit (waktu bebas) untuk dosis kumulatif Cr VI", "Dekontaminasi dan suportif. Tidak ada antidotum spesifik; BAL terbatas; DMSA/CaNa2EDTA dipertimbangkan pada fase awal ingesti"],
        ["**Merkuri (Hg)**", "Tremor, eretisme, gingivostomatitis, nefropati, akrodinia; MeHg: gangguan sensorik/cerebellar (Minamata)", "Hg anorganik urin, sebelum shift: 20 µg/g kreatinin (BEI ACGIH dan IPB Permenaker); >100 µg/g kreatinin = medical removal + evaluasi kelasi. Hg darah/rambut untuk metilmerkuri", "DMSA (oral) atau DMPS; **BAL dikontraindikasikan pada Hg elemental dan metilmerkuri**"],
        ["**Arsen (As)**", "Akut: gastroenteritis berat, Mees lines; kronis: hiperkeratosis, hiperpigmentasi, neuropati, kanker kulit/paru/kandung kemih; arsin: hemolisis", "As anorganik + MMA + DMA urin, akhir shift akhir minggu kerja, puasa seafood 48 sampai 72 jam. ACGIH 2026 15 µg/g kreatinin (Pop); IPB Permenaker 35 µg As/L", "BAL/DMPS/DMSA pada keracunan arsen. **Arsin: transfusi tukar, suportif; kelasi tidak bermanfaat**"],
        ["**Timbal (Pb)**", "Kolik abdomen, anemia (basophilic stippling), neuropati motorik, ensefalopati, nefropati, hipertensi, garis Burton", "BLL (darah): BEI 200 µg/L (20 µg/dL); ZPP sebagai pelengkap (ZPP tinggi dengan BLL rendah = pajanan lama/kelasi diam-diam/defisiensi besi)", "Kelasi (CaNa2EDTA, DMSA/succimer) sesuai BLL dan gejala: BLL 50 sampai 79 µg/dL hanya bila bergejala; ≥80 sampai 100 atau ensefalopati = rawat inap dan kelasi agresif. Perhatikan urutan pada ensefalopati (lihat 3A.4)"],
        ["**Kadmium (Cd)**", "Nefropati tubular proksimal, proteinuria molekul rendah (β2-MG), osteomalasia/itai-itai, emfisema, kanker paru; akut: pneumonitis kimia (laten 4 sampai 12 jam)", "Cd urin: 5 µg/g kreatinin (beban tubuh); Cd darah: 5 µg/L (pajanan terkini); β2-mikroglobulin urin sebagai efek", "**Kelasi (BAL, EDTA) dikontraindikasikan** (memperberat nefrotoksisitas). Akut: suportif, kortikosteroid dipertimbangkan"],
        ["**Mangan (Mn)**", "Manganisme: fase awal psikiatrik (locura manganica), lanjut ekstrapiramidal; pembeda dari Parkinson (respons levodopa buruk)", "**Tidak ada BEI ACGIH.** Mn darah (rujukan 4 sampai 15 µg/L), MRI T1 pallidal", "Hentikan pajanan; CaNa2EDTA efektivitas terbatas karena sulit menembus sawar darah otak; Na-PAS dilaporkan memperbaiki klinis"],
        ["**Nikel (Ni)**", "Dermatitis kontak alergi, asma/rinitis, perforasi septum, kanker paru dan sinonasal; nikel karbonil: pneumonitis tertunda", "Ni urin/serum. Notebook mencantumkan nilai BEI 30 µg/L dan 5 µg/L (spesimen: lihat 3A.7; verifikasi pada buku BEI ACGIH 2026)", "Nikel karbonil: dietilditiokarbamat (Na-DDC, ditiokarb); CaNa2EDTA tidak mencegah kematian paru. Dermatitis: hindari kontak"],
        ["**Berilium (Be)**", "Akut: pneumonitis kimia, dermatitis; kronis: sensitisasi (BeS) dan chronic beryllium disease (granuloma non-kaseosa), kanker paru", "**Tidak ada BEI.** BeLPT darah (berulang) untuk sensitisasi; BAL-BeLPT/biopsi untuk CBD", "Kortikosteroid sistemik untuk CBD bergejala/fungsi paru memburuk; hentikan pajanan; **tidak ada antidotum, kelasi tidak efektif**"],
        ["**Kobalt / hard metal**", "Asma okupasi, hard metal lung disease (GIP, fibrosis), kardiomiopati, dermatitis", "Co urin: BEI 15 µg/L (Co darah: Nq)", "Hentikan pajanan (medical removal), bronkodilator/kortikosteroid; **kelasi tidak efektif**"],
        ["**Seng (Zn)**", "Metal fume fever: demam, menggigil, mialgia, leukositosis, onset 4 sampai 12 jam, sembuh sendiri, tachyphylaxis", "Tidak ada BEI", "Suportif/simtomatik; **kelasi tidak diindikasikan**"],
        ["**Talium (Tl)**", "Nyeri neuropati, GI; alopesia 2 sampai 3 minggu, neuropati ascending, garis Mees", "Tl urin; tidak ada BEI (lihat 3A.11)", "Prussian blue oral, arang aktif berulang, hemodialisis; **hindari Na-DDC** dan hati-hati mobilisasi kalium"],
    ]
    M.add_table(doc, rows, font=8, widths=[2.4, 4.6, 5.4, 4.6])


def penutup(doc):
    doc.add_heading("6. BAGIAN PENUTUP (3A. LOGAM BERAT)", level=1)
    doc.add_heading("6.1 Lafaz OSCE: satu kalimat per zat prioritas", level=2)
    para(doc, "Disusun dari ringkasan bagian 3A dan Quick Table. Gunakan sebagai kerangka jawaban lisan.", size=9.5, italic=True)
    for t in [
        "**Timbal:** \"Pekerja ini keracunan timbal dengan BLL berapa pun yang bergejala aktif; tarik dari pajanan, ukur BLL dan ZPP, dan kelasi sesuai ambang BLL dan gejala, dengan BEI 200 µg/L sebagai acuan pemantauan.\"",
        "**Merkuri:** \"Hg anorganik urin sebelum shift dibandingkan BEI 20 µg/g kreatinin; pada Hg elemental dan metilmerkuri BAL dikontraindikasikan, gunakan DMSA atau DMPS.\"",
        "**Arsen:** \"Pemeriksaan urin harus dispesiasi dan pekerja puasa seafood 48 sampai 72 jam, karena arsenobetain memberi hasil semu tinggi; arsin ditangani dengan transfusi tukar, bukan kelasi.\"",
        "**Kadmium:** \"Nefropati tubular ditandai β2-mikroglobulin urin; kelasi dengan BAL atau EDTA dikontraindikasikan.\"",
        "**Kromium VI:** \"Septum nasi berulserasi dan kanker paru pada pekerja pelapisan krom; Cr VI karsinogen Grup 1, dan NAB Permenaker 5/2018 lebih tua daripada TLV ACGIH terbaru.\"",
        "**Mangan:** \"Gejala ekstrapiramidal pada pengelas tanpa respons levodopa mengarah ke manganisme, bukan Parkinson; tidak ada BEI ACGIH sehingga diagnosis bertumpu pada riwayat pajanan, klinis, dan MRI T1.\"",
        "**Nikel:** \"Dermatitis kontak alergi dan kanker sinonasal pada pemurnian nikel; keracunan nikel karbonil memerlukan observasi panjang karena gejala tertunda, dengan dietilditiokarbamat sebagai antidotum.\"",
        "**Berilium:** \"BeLPT positif berulang dengan granuloma non-kaseosa adalah penyakit berilium kronik; tidak ada BEI dan tidak ada antidotum, terapi utamanya kortikosteroid dan penghentian pajanan.\"",
        "**Kobalt/hard metal:** \"Hard metal lung disease dengan giant cell interstitial pneumonia memerlukan penghentian pajanan dan kortikosteroid; kelasi tidak berguna.\"",
        "**Seng:** \"Demam dan menggigil 4 sampai 12 jam setelah pengelasan galvanis adalah metal fume fever, sembuh sendiri dalam sekitar 24 jam, tidak dikelasi.\"",
        "**Talium:** \"Nyeri neuropati disusul rambut rontok 2 sampai 3 minggu mengarah ke talium; berikan Prussian blue dan arang aktif, hindari ditiokarb.\"",
    ]:
        bullets(doc, [t])

    doc.add_heading("6.2 Vignette MCQ dan jawaban (3A)", level=2)
    para(doc, "Skenario fiktif untuk latihan. Jawaban mengikuti isi bagian 3A di atas.", size=9.5, italic=True)
    V = [
        ("Vignette 1 (Timbal)",
         "Pekerja pabrik aki 41 tahun mengeluh kolik abdomen berat dan lemah pergelangan tangan sejak 2 minggu. Hb 9,1 g/dL dengan basophilic stippling; BLL 62 µg/dL. Tidak ada tanda ensefalopati. Langkah paling tepat?",
         ["A. Observasi dan ulang BLL 3 bulan", "B. Tarik dari pajanan dan beri terapi kelasi karena BLL 50 sampai 79 µg/dL dengan gejala aktif", "C. Berikan BAL tunggal", "D. Cukup suplementasi besi", "E. Kembali bekerja dengan APD"],
         "B. Kelasi pada BLL 50 sampai 79 µg/dL dipertimbangkan bila ada gejala aktif, setelah pekerja ditarik dari pajanan (3A.4). A salah karena gejala aktif; C tidak dipakai tunggal; D dan E mengabaikan pajanan."),
        ("Vignette 2 (Kadmium)",
         "Pekerja baterai Ni-Cd 52 tahun dengan proteinuria. β2-mikroglobulin urin meningkat, Cd urin 14 µg/g kreatinin (BEI 5). Terapi yang dikontraindikasikan?",
         ["A. Penghentian pajanan", "B. Kelasi BAL atau CaNa2EDTA", "C. Hidrasi dan pemantauan fungsi ginjal", "D. Suplementasi kalsium dan vitamin D bila ada osteomalasia", "E. Rujuk nefrologi"],
         "B. Kelasi BAL/EDTA pada nefropati kadmium dikontraindikasikan karena kompleks logam-kelator memperberat kerusakan tubulus (3A.5)."),
        ("Vignette 3 (Berilium)",
         "Pekerja pengecoran paduan Cu-Be 38 tahun, batuk dan sesak progresif. BeLPT darah positif dua kali; biopsi transbronkial menunjukkan granuloma non-kaseosa. Pernyataan yang benar?",
         ["A. BEI berilium urin digunakan untuk diagnosis", "B. Kelasi CaNa2EDTA adalah terapi utama", "C. Ini penyakit berilium kronik; terapi kortikosteroid sistemik dan penghentian pajanan", "D. Kembali bekerja setelah steroid", "E. Diagnosis sarkoidosis cukup"],
         "C. CBD adalah hipersensitivitas tipe IV; tidak ada BEI dan tidak ada antidotum; kortikosteroid dipakai pada penyakit bergejala atau fungsi paru memburuk (3A.8). E keliru karena BeLPT positif dan riwayat pajanan membedakan CBD dari sarkoidosis."),
        ("Vignette 4 (Arsen, jebakan seafood)",
         "Pekerja peleburan tembaga: arsen urin total 480 µg/L pada sampel pagi setelah makan kerang semalam, tanpa gejala. Langkah berikut yang paling tepat?",
         ["A. Diagnosis keracunan arsen dan kelasi segera", "B. Ulang pemeriksaan dengan spesiasi (iAs + MMA + DMA) setelah puasa seafood 48 sampai 72 jam", "C. Abaikan karena asimtomatik", "D. Transfusi tukar", "E. Periksa kadar timbal"],
         "B. Arsenik total urin terdistorsi arsenobetain dari makanan laut; sampel yang benar adalah spesies anorganik ditambah metabolit termetilasi, akhir shift akhir minggu kerja (3A.3)."),
    ]
    for title, stem, opts, ans in V:
        para(doc, "**" + title + "**", after=2)
        para(doc, stem, after=2)
        for o in opts:
            para(doc, o, after=0, style=None).paragraph_format.left_indent = Cm(0.8)
        para(doc, "**Jawaban:** " + ans, size=9.5, after=8)

    doc.add_heading("6.3 Vignette hitung ARKL (data hipotetis)", level=2)
    callout(doc, "Semua angka di bawah ini hipotetis untuk latihan; nilai RfD dan CSF bukan nilai resmi suatu zat. Rumus mengikuti bagian 2.4 dan bagian 4.", fill="E2EFDA")
    para(doc, "**A. Non-kanker.** Pekerja terpajan debu logam X di udara 0,02 mg/m3. Laju inhalasi R = 0,83 m3/jam, tE = 8 jam/hari, fE = 250 hari/tahun, Dt = 10 tahun, Wb = 70 kg, tavg = 10 × 365 = 3.650 hari, RfD inhalasi hipotetis = 1 × 10⁻³ mg/kg/hari.", after=2)
    for t in [
        "Intake = (0,02 × 0,83 × 8 × 250 × 10) / (70 × 3.650) = 332 / 255.500 = 1,30 × 10⁻³ mg/kg/hari.",
        "RQ = 1,30 × 10⁻³ / 1 × 10⁻³ = 1,30 → RQ > 1, risiko tidak aman.",
        "Konsentrasi aman = (RfD × Wb × tavg) / (R × tE × fE × Dt) = (10⁻³ × 255.500) / 16.600 = 0,0154 mg/m3. Konsentrasi saat ini harus diturunkan sekitar 23%.",
    ]:
        bullets(doc, [t])
    para(doc, "**B. Kanker.** Karsinogen Y di udara 0,0005 mg/m3; R = 0,83 m3/jam, tE = 8 jam/hari, fE = 250 hari/tahun, Dt = 30 tahun, Wb = 70 kg, tavg = 70 × 365 = 25.550 hari, CSF hipotetis = 1,5 (mg/kg/hari)⁻¹.", after=2)
    for t in [
        "Intake karsinogenik = (0,0005 × 0,83 × 8 × 250 × 30) / (70 × 25.550) = 24,9 / 1.788.500 = 1,39 × 10⁻⁵ mg/kg/hari.",
        "ECR = 1,39 × 10⁻⁵ × 1,5 = 2,1 × 10⁻⁵.",
        "Interpretasi: di atas 1 × 10⁻⁶ tetapi di bawah 1 × 10⁻⁴. Menurut batas Pedoman ARKL (1E-4) risiko masih dapat diterima, namun di atas target 1E-6 sehingga tetap perlu pengendalian dan pemantauan (prinsip ALARA).",
    ]:
        bullets(doc, [t])


def quality(doc):
    doc.add_heading("7. CATATAN KUALITAS DAN HAL YANG PERLU DICEK", level=1)
    para(doc, "Daftar ini dibuat agar jawaban yang kurang pasti tidak diterima begitu saja.", size=9.5, italic=True)
    bullets(doc, [
        "**PP 88/2019 sudah dicabut** oleh PP 28/2024 (Pasal 1170) menurut NotebookLM; teks utuh PP 88/2019 tidak ada di notebook. Outline yang menyebut PP 88/2019 sebagai regulasi berlaku perlu diperbarui. Cek pasal pada PDF PP 28/2024.",
        "**Dua versi BEI kromium:** 25 µg/L (akhir shift akhir minggu kerja; dipakai IPB Permenaker 5/2018) dan 0,7 µg/L berbasis populasi (ACGIH 2026). Pastikan mana yang ditanyakan dosen.",
        "**Nikel:** notebook menyebut nilai BEI 30 µg/L dan 5 µg/L. Spesimen dan status BEI nikel perlu diverifikasi pada buku TLV/BEI ACGIH 2026 sebelum dihafal.",
        "**Angka Permenaker 5/2018 vs ACGIH:** banyak NAB/IPB Permenaker lebih tua daripada ACGIH 2026 (kromium, arsen BEI 35 vs 15 µg/g, dan lain-lain). Sebutkan keduanya.",
        "**Sumber `CASARE~1.PDF`** di notebook gagal terunggah (error), sehingga tidak dipakai.",
        "**Sumber tambahan 3A** (ATSDR ToxProfile, EPA IRIS, ATSDR CSEM) belum ditambahkan ke notebook. Rekomendasi: cukup CSEM per logam bila diperlukan, karena ToxProfile 300 sampai 600 halaman per zat terlalu berat. Tautan ToxProfile yang sudah diverifikasi: tp2 (As), tp5 (Cd, draf), tp7 (Cr), tp13 (Pb), tp151 (Mn), tp46 (Hg), tp15 (Ni), tp4 (Be), tp33 (Co), tp60 (Zn), tp54 (Tl, draf), pada `atsdr.cdc.gov/ToxProfiles/tpN.pdf`.",
        "**Belum ada:** Permenkes 7/2019 dan regulasi BPJS (bagian 2.5), grafik/gambar (kurva dosis-respons, jalur metabolisme), serta bagian 3B sampai 3F dan quick table zat lain (notebook lain).",
        "**Konten Seng dan Talium** dijawab dalam satu pertanyaan NotebookLM (lebih ringkas, sesuai permintaan memperpendek zat minor). Kobalt juga satu pertanyaan.",
        "**Sumber jawaban** adalah keluaran NotebookLM yang dapat keliru ('can make mistakes'). Angka yang dipakai untuk ujian sebaiknya dicek ke sumber primer.",
    ], size=10)
