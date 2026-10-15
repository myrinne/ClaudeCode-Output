"""Convert NotebookLM markdown answers into python-docx content (helpers)."""
import re
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SUP = str.maketrans("0123456789+-=()ni", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿⁱ")
SUB = str.maketrans("0123456789+-=()", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎")
GREEK = {
    "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "Delta": "Δ", "mu": "µ",
    "lambda": "λ", "sigma": "σ", "Sigma": "Σ", "pi": "π", "theta": "θ", "epsilon": "ε",
    "tau": "τ", "omega": "ω", "Omega": "Ω", "rho": "ρ", "eta": "η", "kappa": "κ",
}
SYMS = {
    "times": "×", "cdot": "·", "le": "≤", "leq": "≤", "ge": "≥", "geq": "≥", "to": "→",
    "rightarrow": "→", "Rightarrow": "⇒", "leftarrow": "←", "approx": "≈", "pm": "±",
    "neq": "≠", "ne": "≠", "infty": "∞", "sum": "Σ", "gg": "≫", "ll": "≪", "div": "÷",
    "degree": "°", "circ": "°", "propto": "∝", "ldots": "…", "dots": "…", "cdots": "…",
    "uparrow": "↑", "downarrow": "↓", "log": "log", "ln": "ln", "min": "min", "max": "max",
    "sim": "~", "equiv": "≡", "lt": "<", "gt": ">",
    "prod": "Π", "implies": "⇒", "quad": " ", "qquad": "  ", "xrightarrow": "→", "bullet": "•",
    "Psi": "Ψ", "AA": "Å",
}


def _sup(m):
    s = m.group(1)
    t = s.translate(SUP)
    return t if all(c in "0123456789+-=()ni⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿⁱ" for c in s) else "^(" + s + ")"


def _sub(m):
    s = m.group(1)
    if all(c in "0123456789+-=()" for c in s):
        return s.translate(SUB)
    return "_" + s if len(s) == 1 else "_(" + s + ")"


def _group(s, i):
    """s[i] == '{'; return (content, index after matching '}')."""
    depth, j = 0, i
    while j < len(s):
        if s[j] == "{":
            depth += 1
        elif s[j] == "}":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1
    return s[i + 1:], len(s)


WRAP = re.compile(r"\\(text|mathrm|mathbf|textbf|mathit|operatorname|textit|boldsymbol)\s*\{")
FRAC = re.compile(r"\\d?frac\s*\{")


def _expand(s):
    """Resolve wrapper commands and fractions with balanced braces (recursive)."""
    out, i = [], 0
    while i < len(s):
        m1, m2 = WRAP.search(s, i), FRAC.search(s, i)
        cands = [m for m in (m1, m2) if m]
        if not cands:
            out.append(s[i:])
            break
        m = min(cands, key=lambda x: x.start())
        out.append(s[i:m.start()])
        brace = m.end() - 1
        a, j = _group(s, brace)
        if m is m1:
            out.append(_expand(a))
            i = j
        else:
            k = j
            while k < len(s) and s[k] in " \t":
                k += 1
            if k < len(s) and s[k] == "{":
                b, j2 = _group(s, k)
                out.append("(" + _expand(a) + ")/(" + _expand(b) + ")")
                i = j2
            else:
                out.append(_expand(a))
                i = j
    return "".join(out)


def latex(s: str) -> str:
    s = s.strip().replace("{,}", ",")
    s = re.sub(r"\\([A-Za-z]+)(?![A-Za-z])",
               lambda m: GREEK.get(m.group(1), SYMS.get(m.group(1), m.group(0))), s)
    s = _expand(s)
    for _ in range(3):
        s = re.sub(r"\^\{([^{}]*)\}", _sup, s)
        s = re.sub(r"_\{([^{}]*)\}", _sub, s)
    for _ in range(6):
        s = re.sub(r"\\d?frac\s*\{([^{}]*)\}\s*\{([^{}]*)\}", r"(\1)/(\2)", s)
    s = re.sub(r"\\sqrt\s*\{([^{}]*)\}", r"√(\1)", s)
    s = re.sub(r"\\(left|right|big|Big|bigg)\s*", "", s)
    s = re.sub(r"\\(begin|end)\{[^}]*\}", "", s)
    s = s.replace("\\\\", " ").replace("\\%", "%").replace("\\_", "_").replace("\\&", "&")
    s = s.replace("\\,", " ").replace("\\;", " ").replace("\\:", " ").replace("\\!", "").replace("\\ ", " ")
    s = re.sub(r"\\([A-Za-z]+)", lambda m: GREEK.get(m.group(1), SYMS.get(m.group(1), m.group(1))), s)
    for _ in range(3):
        s = re.sub(r"\^\{([^{}]*)\}", _sup, s)
        s = re.sub(r"_\{([^{}]*)\}", _sub, s)
    s = re.sub(r"\^([0-9+\-])", lambda m: m.group(1).translate(SUP), s)
    s = re.sub(r"_([0-9])", lambda m: m.group(1).translate(SUB), s)
    s = s.replace("{", "").replace("}", "").replace("&", " ")
    return re.sub(r"[ \t]+", " ", s).strip()


DISP = re.compile(r"\\{1,2}\[(.+?)\\{1,2}\]", re.S)
INL = re.compile(r"\\{1,2}\((.+?)\\{1,2}\)", re.S)


def demath(text: str) -> str:
    text = INL.sub(lambda m: latex(m.group(1)), text)
    text = DISP.sub(lambda m: latex(m.group(1)), text)
    return text


TOK = re.compile(r"(\*\*.+?\*\*|(?<![\w*])\*(?!\s).+?(?<!\s)\*(?![\w*])|`[^`]+`)")


def add_inline(par, text, size=None, base_bold=False, base_italic=False, color=None):
    text = demath(text).replace("<br>", "\n").replace("<br/>", "\n")
    text = re.sub(r"\[\s*(\d+(?:\s*,\s*\d+)*)\s*\]", "", text)  # stray citation numbers
    def emit(t, bold, italic):
        parts = t.split("\n")
        for i, p in enumerate(parts):
            if i:
                par.add_run().add_break()
            if not p:
                continue
            r = par.add_run(p)
            r.bold = bold or None
            r.italic = italic or None
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor(*color)
    def walk(t, bold, italic):
        pos = 0
        for m in TOK.finditer(t):
            if m.start() > pos:
                emit(t[pos:m.start()], bold, italic)
            tok = m.group(0)
            if tok.startswith("**"):
                walk(tok[2:-2], True, italic)
            elif tok.startswith("`"):
                emit(tok[1:-1], bold, italic)
            else:
                walk(tok[1:-1], bold, True)
            pos = m.end()
        if pos < len(t):
            emit(t[pos:], bold, italic)
    walk(text, base_bold, base_italic)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def add_table(doc, rows, header=True, font=8.5, widths=None):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    for ri, r in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(ncol):
            c = cells[ci]
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            txt = r[ci] if ci < len(r) else ""
            add_inline(p, txt, size=font, base_bold=(header and ri == 0))
            if header and ri == 0:
                shade(c, "DCE6F1")
    if widths:
        for row in t.rows:
            for ci, w in enumerate(widths):
                row.cells[ci].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def clean_answer(md: str):
    """Return (body_lines, refs_lines)."""
    lines = md.replace("\r", "").split("\n")
    # drop preamble before first '---'
    for i, l in enumerate(lines[:40]):
        if l.strip() == "---":
            lines = lines[i + 1:]
            break
    body, refs, mode = [], [], "body"
    for l in lines:
        s = l.strip()
        if mode == "body" and re.match(r"^#{1,4}\s*.*(Sumber|Rujukan|Referensi|Daftar Pustaka)", s, re.I) and not s.startswith("| "):
            mode = "refs"
            continue
        if mode == "refs":
            if s == "" or s == "---" and not refs:
                continue
            if re.match(r"^(\d+[.)]|[*\-•])\s", s):
                refs.append(s)
                continue
            mode = "tail"
            continue
        if mode == "tail":
            continue
        if "🩺" in s:
            continue
        body.append(l)
    return body, refs


def render(doc, body, h_off=0, skip_h_prefix=None):
    """Render markdown body lines to doc. h_off: added to heading level (### -> H(3+off-1))"""
    i = 0
    n = len(body)
    while i < n:
        l = body[i].rstrip()
        s = l.strip()
        if s == "" or s == "---" or s == "***":
            i += 1
            continue
        # fenced code / ASCII diagram -> monospace block
        if s.startswith("```"):
            i += 1
            block = []
            while i < n and not body[i].strip().startswith("```"):
                block.append(body[i].rstrip("\n").rstrip())
                i += 1
            i += 1
            if any(x.strip() for x in block):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.3)
                p.paragraph_format.space_after = Pt(6)
                for k, line in enumerate(block):
                    if k:
                        p.add_run().add_break()
                    r = p.add_run(line.replace("\t", "    "))
                    r.font.name = "Consolas"
                    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
                    r.font.size = Pt(6.5)
            continue
        # display math block
        if s.startswith("\\\\[") or s.startswith("\\["):
            buf = [s]
            while not re.search(r"\\{1,2}\]\s*$", buf[-1]) and i + 1 < n:
                i += 1
                buf.append(body[i].strip())
            txt = demath(" ".join(buf))
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, txt, base_italic=True)
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            lvl = len(m.group(1))
            text = m.group(2)
            txt = re.sub(r"\*+", "", demath(text)).strip()
            if skip_h_prefix and txt.upper().startswith(skip_h_prefix):
                i += 1
                continue
            level = {1: 2, 2: 2, 3: 2, 4: 3}[lvl] if h_off == 0 else min({1: 2, 2: 2, 3: 2, 4: 3}[lvl] + h_off, 4)
            doc.add_heading(txt, level=level)
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < n and body[i].strip().startswith("|"):
                row = body[i].strip()
                i += 1
                if re.match(r"^\|\s*:?-{2,}", row):
                    continue
                cells = [c.strip() for c in row.strip("|").split("|")]
                rows.append(cells)
            if rows:
                nc = len(rows[0])
                if nc == 2:
                    add_table(doc, rows, header=True, font=9, widths=[3.0, 14.2])
                else:
                    add_table(doc, rows, header=True, font=8 if nc >= 4 else 9)
            continue
        m = re.match(r"^(\s*)([*\-•]|\d+[.)])\s+(.*)$", l)
        if m:
            indent = len(m.group(1).replace("\t", "    "))
            marker, text = m.group(2), m.group(3)
            if marker[0].isdigit():
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6 + 0.6 * (indent // 4))
                p.paragraph_format.first_line_indent = Cm(-0.5)
                p.paragraph_format.space_after = Pt(2)
                add_inline(p, marker + " " + text)
            else:
                style = "List Bullet" if indent < 2 else ("List Bullet 2" if indent < 6 else "List Bullet 3")
                p = doc.add_paragraph(style=style)
                p.paragraph_format.space_after = Pt(1)
                add_inline(p, text)
            i += 1
            continue
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            add_inline(p, s.lstrip("> "), base_italic=True)
            i += 1
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        add_inline(p, s)
        i += 1


def render_refs(doc, refs, label="Sumber yang dirujuk NotebookLM"):
    if not refs:
        return
    p = doc.add_paragraph()
    r = p.add_run(label + ":")
    r.bold = True
    r.font.size = Pt(8)
    for ref in refs:
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Cm(0.5)
        q.paragraph_format.space_after = Pt(0)
        add_inline(q, re.sub(r"^(\d+[.)]|[*\-•])\s+", "", ref), size=8, color=(0x59, 0x59, 0x59))
