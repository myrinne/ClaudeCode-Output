# -*- coding: utf-8 -*-
"""Tambah bagian 6.0 'Lafaz OSCE berbentuk kasus' (3 vignette: timbal, benzena, organofosfat) tepat sebelum
6.1. Bagian 6.1 dst (lafaz satu kalimat per zat) tidak diubah.

Kasus lain (arsin, merkuri, metanol, n-heksana, paraquat, CO, H2S, HF, silika, asbes, kapas) tersimpan di
kasus_lafaz.py sebagai cadangan; tambahkan ke PILIHAN bila diminta.

Pemakaian: python rewrite_lafaz.py <docx_in> <docx_out>
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
import mdconv as M
from kasus_lafaz import SECTIONS

# (bagian, indeks kasus, label zat)
PILIHAN = [("6.1", 0, "logam berat"), ("6.4", 0, "pelarut organik"), ("6.6", 0, "pestisida")]

SRC, DST = sys.argv[1], sys.argv[2]
doc = Document(SRC)
body = doc.element.body


def text_of(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


kids = [k for k in body.iterchildren() if k.tag != qn("w:sectPr")]
anchor = next(k for k in kids if k.tag == qn("w:p") and text_of(k).startswith("6.1 Lafaz OSCE"))

n_before = len(kids)
doc.add_heading("6.0 Lafaz OSCE berbentuk kasus (vignette)", level=2)
p = doc.add_paragraph()
r = p.add_run("Skenario fiktif untuk latihan. Tiap kasus berisi tanda keracunan, lalu tiga pertanyaan penguji: "
              "dugaan diagnosis, tatalaksana, serta surveilans dan tindak lanjut okupasi. Lafaz diucapkan ke penguji. "
              "Lafaz satu kalimat per zat ada di 6.1 sampai 6.10.")
r.italic = True
for n, (num, idx, label) in enumerate(PILIHAN, 1):
    k = SECTIONS[num]["kasus"][idx]
    judul = "Kasus %d (%s): %s" % (n, label, k["judul"].split(": ", 1)[1])
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(judul); r.bold = True
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run("Vignette: "); r.bold = True
    p.add_run(k["vignette"])
    rows = [["Pertanyaan penguji", "Lafaz OSCE (diucapkan ke penguji)"]]
    for i, (q, a) in enumerate(k["qa"], 1):
        rows.append(["%d. %s" % (i, q), a])
    M.add_table(doc, rows, header=True, font=9, widths=[4.6, 12.6])

new = [k for k in body.iterchildren() if k.tag != qn("w:sectPr")][n_before:]
for el in new:
    anchor.addprevious(el)

doc.save(DST)
print("saved", DST)
