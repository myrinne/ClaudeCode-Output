from docx.shared import Pt, Cm
import mdconv as M
from extras import para, bullets, callout


def regulasi_lain(doc):
    doc.add_heading("2.5 Regulasi yang belum tercakup dan pegangan cepat", level=2)
    bullets(doc, [
        "**Permenkes 7/2019 (kesehatan lingkungan rumah sakit)** dan **regulasi BPJS Ketenagakerjaan untuk PAK** tidak ada di notebook toksikologi logam berat; perlu diisi dari sumber lain.",
        "**Pegangan cepat:** NAB dan IPB = Permenaker 5/2018; daftar PAK dan hak JKK sampai 3 tahun setelah hubungan kerja berakhir = Perpres 7/2019; pelayanan PAK = Permenkes 11/2022; kesehatan kerja = PP 28/2024 (Pasal 232 sampai 244); ARKL = Pedoman Ditjen PP&PL Kemenkes 2012.",
    ])


def template_intro(doc):
    para(doc, "Setiap zat disajikan sebagai satu tabel berisi 13 poin template, lalu **pearl dan jebakan ujian** (poin 14), dan **infografik satu halaman** (kunci hafalan: sumber, karakteristik, efek akut dan kronik, biomarker, terapi, clinical pearl). Poin 15 (tabel singkat lintas zat) ada di bagian 5.", size=10)
    M.add_table(doc, [
        ["No", "Poin template"],
        ["1", "Sifat fisik dan kimia / klasifikasi"],
        ["2", "Sumber (penggunaan industri)"],
        ["3", "Toksikokinetik"],
        ["4", "Toksikodinamik"],
        ["5", "Biomonitoring: biomarker dan waktu sampling"],
        ["6", "Waktu paruh (T1/2)"],
        ["7", "Efek biologis akut"],
        ["8", "Efek biologis kronis"],
        ["9", "Standar exposure limit (TLV, NAB)"],
        ["10", "Standar BEI / IPB"],
        ["11", "Tatalaksana akut / antidotum"],
        ["12", "Kelaikan kerja dan manajemen (individu dan komunitas)"],
        ["13", "Keterangan lain (karsinogenitas, dst.)"],
        ["14", "Pearl dan jebakan ujian (di bawah tabel)"],
        ["15", "Tabel singkat lintas zat (bagian 5A dan 5B)"],
    ], font=9, widths=[1.2, 15.6])
    callout(doc,
            "**Jebakan lintas zat.** Angka NAB dan IPB Permenaker 5/2018 sering lebih tua daripada TLV/BEI ACGIH 2026 (contoh: Cr VI larut air, "
            "NAB 0,05 mg/m3 vs TLV 0,0002 mg/m3). Sebutkan keduanya dan nyatakan acuan yuridis di Indonesia. Diagnosis PAK tidak dirinci sebagai urutan langkah; "
            "kausalitas dinilai secara utuh (klinis, pajanan, dosis, faktor pejamu dan luar kerja).",
            fill="FDE9D9")


def formula(doc):
    doc.add_heading("4. FORMULA WAJIB HITUNG", level=1)
    para(doc, "Rumus yang bukan berasal dari NotebookLM diberi tanda (baku).", size=9.5, italic=True)
    M.add_table(doc, [
        ["Kelompok", "Rumus", "Keterangan"],
        ["TWA 8 jam", "TWA = Σ(Ci × Ti) / ΣTi", "Contoh: 35 ppm × 2 j, 15 × 3, 25 × 1, 5 × 2 → 150/8 = 18,75 ppm (1.4)."],
        ["Adjusted OEL shift panjang (Brief & Scala)", "RFd = (8/hd) × ((24 − hd)/16); RFw = (40/hw) × ((168 − hw)/128); OELadj = OEL8 jam × RF (RF terkecil)", "Shift 12 jam: RFd = 0,50 → NAB 100 ppm menjadi 50 ppm (1.4)."],
        ["Campuran (unity)", "Σ(Ci / TLVi) ≤ 1 untuk efek/organ target sama", "Efek independen dinilai terpisah (1.5)."],
        ["Non-kanker (ARKL)", "Intake = (C × R × tE × fE × Dt) / (Wb × tavg); tavg = Dt × 365; RQ = Intake / RfD; inhalasi RQ = C / RfC", "RQ ≤ 1 aman; RQ > 1 tidak aman (2.4)."],
        ["Konsentrasi aman (ARKL)", "C aman = (RfD × Wb × tavg) / (R × tE × fE × Dt)", "(baku) turunan aljabar."],
        ["Kanker", "tavg = 70 × 365 = 25.550 hari; ECR = Intake × CSF (atau C × IUR)", "ECR ≤ 1E-4 dapat diterima; target 1E-6 (2.4)."],
        ["MOE", "MOE = NOAEL (atau BMDL) / pajanan", "(baku) makin besar makin aman."],
        ["Koreksi kreatinin", "µg/g kreatinin = konsentrasi urin (µg/L) / kreatinin (g/L)", "Diterima bila kreatinin 0,3 sampai 3 g/L (1.8)."],
        ["Koreksi berat jenis", "C koreksi = C ukur × (SG ref − 1,000) / (SG sampel − 1,000)", "SG ref 1,024; sampel layak 1,010 sampai 1,030. Contoh: 10 µg/L pada SG 1,012 → 20 µg/L (1.8)."],
        ["Steady state", "≈ 5 × T1/2 (± 97% Css)", "(1.10)"],
        ["Kinetik", "T1/2 = 0,693 / Kel = 0,693 × Vd / Cl; sisa setelah n × T1/2 = (1/2)^n; faktor akumulasi = 1 / (1 − e^(−Kel × τ))", "(1.10); faktor akumulasi dan sisa (baku)."],
    ], font=8.5, widths=[3.2, 7.4, 6.2])


def quick_table(doc):
    doc.add_heading("5A. QUICK TABLE AKHIR (3A. LOGAM BERAT)", level=1)
    para(doc, "Kolom: zat | gejala kunci | biomarker (sampel + waktu) | antidotum / terapi kunci. Nilai yang berbeda antar sumber ada di bagian 7.", size=9.5, italic=True)
    rows = [
        ["Zat", "Gejala kunci", "Biomarker (sampel + waktu)", "Antidotum / terapi"],
        ["**Kromium VI**", "Ulkus dan perforasi septum, chrome holes, dermatitis, asma, kanker paru/sinonasal", "Cr total urin akhir shift akhir minggu (IPB 25 µg/L; ACGIH 2026 0,7 µg/L Pop); Cr eritrosit untuk Cr VI", "Suportif; asam askorbat (ingesti, kulit); tidak ada antidotum spesifik"],
        ["**Merkuri**", "Tremor intensi, eretisme, gingivostomatitis, nefropati; MeHg: Minamata", "U-Hg prashift 20 µg/g kreatinin (Hg⁰, anorganik); darah dan rambut untuk MeHg", "DMPS atau DMSA; **BAL dilarang**"],
        ["**Arsen**", "Akut: gastroenteritis, QT, Mees; kronis: raindrop, keratosis, neuropati, kanker; arsin: hemolisis", "iAs + MMA + DMA urin akhir minggu, puasa seafood 48 sampai 72 jam (BEI 15 µg/g; IPB 35 µg/L)", "BAL/DMSA/DMPS (iAs); **arsin: transfusi tukar + alkalinisasi, tanpa kelasi**"],
        ["**Timbal**", "Kolik, anemia + stippling, neuropati motorik, ensefalopati, nefropati, hipertensi", "BLL (BEI 200 µg/L) bebas waktu; ZPP time lag 8 sampai 12 minggu", "DMSA oral; ensefalopati: **BAL dulu, CaNa2EDTA 4 jam kemudian**"],
        ["**Kadmium**", "Tubulopati (β2-MG), Itai-itai, emfisema, kanker paru; fume: pneumonitis laten", "U-Cd 5 µg/g kreatinin (beban tubuh), B-Cd 5 µg/L (terkini); β2-MG / RBP", "**Kelasi dilarang**; suportif"],
        ["**Mangan**", "Locura mangánica lalu ekstrapiramidal (cock-walk, retropulsi); levodopa gagal", "**Tidak ada BEI**; Mn darah 4 sampai 15 µg/L; MRI T1 pallidum", "Hentikan pajanan; CaNa2EDTA/Na-PAS; **BAL dilarang**"],
        ["**Nikel**", "Dermatitis, asma, perforasi septum, kanker paru/sinonasal; Ni(CO)4 bifasik", "Ni urin akhir minggu (BEI menurut NotebookLM 30 dan 5 µg/L; verifikasi)", "Ni(CO)4: **Na-DDC / disulfiram**; observasi 24 sampai 48 jam"],
        ["**Berilium**", "CBD, sensitisasi, ulkus, kanker paru; sarkoidosis tiruan", "**Tidak ada BEI**; BeLPT (darah, BAL)", "Prednison (CBD bergejala); BeS: medical removal; **tanpa kelasi**"],
        ["**Kobalt / hard metal**", "Asma, HMLD/GIP (emperipolesis), kardiomiopati, goiter, polisitemia", "Co urin 15 µg/L (Ns); hard metal: Ns, Nq", "Medical removal + prednison; **tanpa kelasi**"],
        ["**Seng**", "Metal fume fever 4 sampai 12 jam, sembuh 24 sampai 48 jam, pola Senin", "Tidak ada BEI", "Suportif; **tanpa antibiotik dan tanpa kelasi**"],
        ["**Talium**", "Kolik + obstipasi, burning soles, alopesia 2 sampai 3 minggu, Mees", "Tl urin 24 jam; tidak ada BEI", "**Prussian blue**, arang aktif; **DDC dilarang**; KCl hati-hati"],
    ]
    M.add_table(doc, rows, font=8, widths=[2.3, 5.0, 5.2, 4.3])


def penutup(doc):
    doc.add_heading("6. BAGIAN PENUTUP (3A DAN 3B)", level=1)
    doc.add_heading("6.1 Lafaz OSCE: satu kalimat per zat (3A)", level=2)
    for t in [
        "**Timbal:** \"Pekerja dengan BLL ≥ 50 µg/dL dan gejala aktif ditarik dari pajanan dan dikelasi (DMSA), sedangkan ensefalopati mendapat BAL lebih dulu lalu CaNa2EDTA 4 jam kemudian; BEI BLL 200 µg/L.\"",
        "**Merkuri:** \"U-Hg prashift dibandingkan BEI 20 µg/g kreatinin; pada Hg elemental dan metilmerkuri BAL dikontraindikasikan, gunakan DMPS atau DMSA, dan metilmerkuri diperiksa dari darah atau rambut.\"",
        "**Arsen:** \"Urin harus dispesiasi dengan puasa seafood 48 sampai 72 jam karena arsenobetain; arsin ditangani dengan alkalinisasi urin dan transfusi tukar tanpa kelasi.\"",
        "**Kadmium:** \"Nefropati tubular ditandai β2-mikroglobulin atau RBP urin (cek pH); kelasi dikontraindikasikan.\"",
        "**Kromium VI:** \"Perforasi septum dan kanker paru pada pekerja plating krom; Cr VI Grup 1, dan NAB Permenaker lebih tua daripada TLV ACGIH 2026.\"",
        "**Mangan:** \"Parkinsonisme tanpa respons levodopa pada pengelas adalah manganisme; tidak ada BEI, diagnosis dari pajanan, klinis, dan MRI T1.\"",
        "**Nikel:** \"Nikel karbonil punya laten 12 sampai 36 jam sehingga observasi 24 sampai 48 jam dengan Na-DDC; dermatitis dikonfirmasi patch test NiSO4 5%.\"",
        "**Berilium:** \"BeLPT positif berulang dengan granuloma non-kaseosa adalah CBD; tidak ada BEI dan antidotum, terapi prednison dan penghentian pajanan.\"",
        "**Kobalt/hard metal:** \"HMLD dengan emperipolesis memerlukan medical removal dan prednison; ketiadaan Co pada biopsi tidak menyingkirkan diagnosis.\"",
        "**Seng:** \"Demam dan menggigil 4 sampai 12 jam setelah las galvanis dengan pola Senin adalah metal fume fever; suportif tanpa antibiotik.\"",
        "**Talium:** \"Kolik, burning soles, lalu alopesia 2 sampai 3 minggu; berikan Prussian blue dan hindari DDC.\"",
    ]:
        bullets(doc, [t])

    doc.add_heading("6.2 Vignette MCQ (3A)", level=2)
    para(doc, "Skenario fiktif untuk latihan.", size=9.5, italic=True)
    V = [
        ("Vignette 1 (Timbal)",
         "Pekerja pabrik aki 41 tahun dengan kolik abdomen dan lemah pergelangan tangan 2 minggu. Hb 9,1 g/dL dengan basophilic stippling; BLL 62 µg/dL; tanpa ensefalopati. Langkah paling tepat?",
         ["A. Observasi, ulang BLL 3 bulan", "B. Tarik dari pajanan dan kelasi (DMSA) karena BLL 50 sampai 79 µg/dL dengan gejala aktif", "C. BAL tunggal", "D. Suplementasi besi saja", "E. Lanjut bekerja dengan APD"],
         "B. Kelasi dipertimbangkan pada BLL 50 sampai 79 dengan gejala aktif setelah pekerja ditarik dari pajanan; DMSA oral lini pertama."),
        ("Vignette 2 (Kadmium)",
         "Pekerja baterai Ni-Cd 52 tahun dengan proteinuria; β2-mikroglobulin urin meningkat, U-Cd 14 µg/g kreatinin (BEI 5). Tindakan yang dikontraindikasikan?",
         ["A. Hentikan pajanan", "B. Kelasi BAL atau CaNa2EDTA", "C. Hidrasi dan pantau ginjal", "D. Kalsium dan vitamin D bila osteomalasia", "E. Rujuk nefrologi"],
         "B. Kelasi memobilisasi Cd ke tubulus dan memicu ATN; terapi suportif dan penghentian pajanan."),
        ("Vignette 3 (Berilium)",
         "Pekerja pemesinan paduan Cu-Be 38 tahun dengan batuk dan sesak progresif; BeLPT darah positif dua kali; biopsi transbronkial: granuloma non-kaseosa. Pernyataan yang benar?",
         ["A. BEI Be urin dipakai untuk diagnosis", "B. CaNa2EDTA terapi utama", "C. CBD: prednison dan penghentian pajanan", "D. Kembali bekerja setelah steroid", "E. Cukup didiagnosis sarkoidosis"],
         "C. CBD adalah hipersensitivitas tipe IV; tidak ada BEI maupun antidotum. BeLPT positif membedakan CBD dari sarkoidosis."),
        ("Vignette 4 (Arsen)",
         "Pekerja smelter tembaga: arsen urin total 480 µg/L pada sampel pagi setelah makan kerang semalam; tanpa gejala. Langkah paling tepat?",
         ["A. Diagnosis keracunan arsen dan kelasi", "B. Ulang dengan spesiasi (iAs + MMA + DMA) setelah puasa seafood 48 sampai 72 jam", "C. Abaikan", "D. Transfusi tukar", "E. Periksa timbal"],
         "B. Arsenobetain dari makanan laut memberi hasil semu tinggi; yang dinilai adalah iAs + metabolit termetilasi akhir shift akhir minggu."),
    ]
    for title, stem, opts, ans in V:
        para(doc, "**" + title + "**", after=2)
        para(doc, stem, after=2)
        for o in opts:
            para(doc, o, after=0).paragraph_format.left_indent = Cm(0.8)
        para(doc, "**Jawaban:** " + ans, size=9.5, after=8)

    doc.add_heading("6.3 Vignette hitung ARKL (data hipotetis)", level=2)
    callout(doc, "Angka hipotetis untuk latihan; RfD dan CSF bukan nilai resmi suatu zat. Default memakai nilai Pedoman ARKL 2012 untuk pekerja: R 0,83 m3/jam, tE 8 jam/hari, fE 250 hari/tahun, Dt 25 tahun, Wb 55 kg.", fill="E2EFDA")
    para(doc, "**A. Non-kanker.** Debu logam X 0,02 mg/m3; Dt = 25 tahun; tavg = 25 × 365 = 9.125 hari; RfD hipotetis = 1 × 10⁻³ mg/kg/hari.", after=2)
    bullets(doc, [
        "Intake = (0,02 × 0,83 × 8 × 250 × 25) / (55 × 9.125) = 830 / 501.875 = 1,65 × 10⁻³ mg/kg/hari.",
        "RQ = 1,65 × 10⁻³ / 1 × 10⁻³ = 1,65 → RQ > 1, tidak aman.",
        "C aman = (RfD × Wb × tavg) / (R × tE × fE × Dt) = (10⁻³ × 501.875) / 41.500 = 0,0121 mg/m3; konsentrasi perlu turun ± 40%.",
    ])
    para(doc, "**B. Kanker.** Karsinogen Y 0,0005 mg/m3; Dt = 25 tahun; tavg = 70 × 365 = 25.550 hari; CSF hipotetis = 1,5 (mg/kg/hari)⁻¹.", after=2)
    bullets(doc, [
        "Intake = (0,0005 × 0,83 × 8 × 250 × 25) / (55 × 25.550) = 20,75 / 1.405.250 = 1,48 × 10⁻⁵ mg/kg/hari.",
        "ECR = 1,48 × 10⁻⁵ × 1,5 = 2,2 × 10⁻⁵.",
        "Di atas target 1E-6 tetapi di bawah 1E-4: masih dapat diterima menurut batas Pedoman ARKL, namun tetap dikendalikan dan dipantau (ALARA).",
    ])


def quality(doc):
    doc.add_heading("7. CATATAN KUALITAS DAN HAL YANG PERLU DICEK", level=1)
    para(doc, "Isi catatan ini berasal dari keluaran NotebookLM yang dapat keliru; angka untuk ujian sebaiknya dicek ke sumber primer.", size=9.5, italic=True)
    bullets(doc, [
        "**PP 88/2019 dicabut** oleh PP 28/2024 (Pasal 1170) menurut NotebookLM; teks PP 88/2019 tidak ada di notebook. Cek pasal pada PDF PP 28/2024.",
        "**BEI kromium ada dua versi:** 25 µg/L (IPB Permenaker; akhir shift akhir minggu) dan 0,7 µg/L berbasis populasi (ACGIH 2026).",
        "**Arsen:** IPB Permenaker 35 µg/L vs ACGIH 2026 15 µg/g kreatinin (Pop).",
        "**Nikel:** nilai BEI 30 dan 5 µg/L belum terverifikasi pada pustaka lain; cek buku BEI ACGIH 2026.",
        "**NAB Permenaker 5/2018 yang berbeda antar keluaran NotebookLM:** mangan (0,2 inhalabel dan 0,02 respirabel pada tabel Lampiran 3 vs 0,1 pada keluaran Bagian B) dan berilium (0,00005 mg/m3 pada tabel vs 0,002 pada keluaran Bagian B). Tabel di bagian 2.1 dipakai sebagai acuan; cek ulang PDF.",
        "**Nomor butir Perpres 7/2019:** keluaran per zat kadang berbeda dari tabel di bagian 2.2 (mis. Pb, Cd, Mn, Ni, Be, Co, Tl, Zn). Tabel bagian 2.2 dipakai sebagai acuan; cek ulang lampiran.",
        "**Belum ada:** Permenkes 7/2019 dan regulasi BPJS, bagian 3C sampai 3F dan quick table zat lain (notebook lain), serta sumber tambahan ATSDR dan EPA.",
        "**Seng, talium, dan kobalt** dijawab ringkas (zat minor).",
        "**Klaim yang perlu diverifikasi:** efek asam askorbat IV dosis tinggi pada kromium (bagian 3A.1).",
        "**3B, angka ACGIH yang berbeda dari pustaka yang saya ketahui (cek buku TLV/BEI 2026):** TLV xilena 20 ppm (pustaka lain 100 ppm, STEL 150, BEI MHA 1,5 g/g); BEI 2,5-HD n-heksana ACGIH 0,5 mg/L (pustaka lain 0,4 mg/L); TLV etilen glikol 25 ppm uap; TLV isopropil eter 20/50 ppm (pustaka lain 250/310); TLV etil asetat 100 ppm (pustaka lain 400); BEI MA+PGA stirena 150 mg/g (ACGIH lama 400); BEI PCE darah 0,4 mg/L.",
        "**3B, nomor butir Perpres 7/2019 dari NotebookLM** (benzena butir 12, derivat nitro/amina 13, keton/alkohol/glikol 14, eter 16, ester 17, halogen alifatik 10, CO 27, CS2 9; Lampiran III butir 4 dan 8) tidak diverifikasi ke lampiran; cek ulang.",
        "**3B, bukti terbatas:** N-asetilsistein untuk CCl4 (rasional pengisian GSH, bukan terapi baku); uji Lampu Wood urin untuk etilen glikol (sensitivitas rendah; catatan tambahan, bukan dari NotebookLM).",
        "**3B, belum ada:** infografik per zat (hanya 3A yang punya), serta vignette hitung dan quick table zat 3C sampai 3F.",
    ], size=10)


def intro_3b(doc):
    para(doc, "Struktur sama dengan 3A: tabel 13 poin per zat lalu pearl dan jebakan ujian (poin 14). Sumber: notebook NotebookLM \"Toxicology - Organic Solvents\" (99 sumber). **Tidak ada infografik untuk 3B.** Poin 12 tidak merinci urutan langkah diagnosis PAK (konsisten dengan 3A). Angka NAB/IPB Permenaker 5/2018 dan TLV/BEI ACGIH yang berbeda ditampilkan berdampingan; yang bertanda \"cek ulang\" ada di bagian 7.", size=10)
    M.add_table(doc, [
        ["Kelompok (outline 3B)", "Zat pada catatan ini"],
        ["Aromatik", "Benzena, toluena, xilena, stirena, anilin dan nitrobenzena (methemoglobinemia)"],
        ["Alifatik", "n-Heksana (neuropati, 2,5-heksanadion)"],
        ["Keton / eter / ester", "MEK dan MBK; aseton, isopropil eter, etil asetat (tabel ringkas)"],
        ["Alkohol / glikol", "Metanol, etilen glikol, PEG"],
        ["Aldehida", "Formaldehida"],
        ["Terhalogenasi", "Vinil klorida monomer, diklorometana, CCl4, trikloroetilen, perkloroetilen"],
        ["Sulfur", "Karbon disulfida"],
        ["Lain", "Dioksin/TCDD dan toluena diisosianat (TDI)"],
    ], font=9, widths=[4.5, 12.3])
    callout(doc,
            "**Jebakan lintas pelarut.** (1) Hampir semua pelarut ini tanpa antidot: terapi suportif, jangan induksi muntah, hindari epinefrin pada hidrokarbon aromatik/terhalogenasi (sensitisasi miokardium). "
            "(2) Sarung tangan lateks hampir selalu salah: pilih Viton/PVA/butil sesuai zat. (3) Biomarker pelarut sering tidak spesifik (Ns) atau berlatar belakang (B): sebutkan notasi dan waktu sampling.",
            fill="FDE9D9")


def quick_table_3b(doc):
    doc.add_heading("5B. QUICK TABLE AKHIR (3B. ORGANIC SOLVENT)", level=1)
    para(doc, "Kolom: zat | gejala kunci | biomarker (sampel + waktu) | antidotum / terapi kunci. Nilai yang ditandai di bagian 7 perlu dicek.", size=9.5, italic=True)
    rows = [
        ["Zat", "Gejala kunci", "Biomarker (sampel + waktu)", "Antidotum / terapi"],
        ["**Benzena**", "Sitopenia, anemia aplastik, MDS/AML; narkosis", "**S-PMA** urin akhir shift (BEI 25 µg/g); t,t-MA 500 µg/g (artefak asam sorbat)", "Tanpa antidot; hindari epinefrin"],
        ["**Toluena**", "SSP, ototoksik + bising, ensefalopati, fetal toluene syndrome", "o-kresol urin 0,3 mg/g (akhir shift); toluena darah 0,02 mg/L (sebelum shift terakhir pekan)", "Suportif; hindari epinefrin"],
        ["**Xilena**", "SSP, ototoksik (p-xilena), kulit", "MHA urin akhir shift (Permenaker 1,5 g/g)", "Suportif"],
        ["**Stirena**", "Diskromatopsia biru-kuning, ototoksik, SSP", "MA + PGA urin akhir shift (Ns)", "Suportif"],
        ["**Anilin / nitrobenzena**", "Sianosis cokelat, Met-Hb, hemolisis tertunda, kandung kemih", "Met-Hb 1,5% Hb; PAP 50 mg/g; PNP 5 mg/g (akhir shift akhir pekan)", "**Biru metilen 1 sampai 2 mg/kg (maks 7)**; G6PD: transfusi tukar"],
        ["**n-Heksana**", "Polineuropati distal, foot drop, coasting", "2,5-HD urin akhir shift tanpa hidrolisis (0,4 mg/L Permenaker)", "Hentikan pajanan, rehabilitasi, AFO"],
        ["**MEK / MBK**", "MEK: iritan; MBK: neuropati; MEK memperkuat n-heksana", "MEK urin 2 mg/L (Ns); 2,5-HD", "Suportif; jangan campur MEK + n-heksana"],
        ["**Metanol**", "Laten 12 sampai 24 jam, Kussmaul, buta salju, nekrosis putamen", "Metanol urin 15 mg/L (okupasi); metanol serum (klinis)", "**Fomepizol**, bikarbonat, folinat, HD"],
        ["**Etilen glikol**", "Mabuk tanpa bau, asidosis, AKI, kristal Ca oksalat", "EG serum, glikolat; sedimen urin", "**Fomepizol**, tiamin + B6, HD; Ca hanya bila simtomatik"],
        ["**Formaldehida**", "Iritasi, asma, DKA, NPC, AML", "Tidak ada BEI; air monitoring", "Suportif; observasi 24 sampai 48 jam"],
        ["**VCM**", "Raynaud, akroosteolisis, sklerodermiform, angiosarkoma hati, trombositopenia", "TDGA urin (Ns); BEI under study", "Tanpa antidot; USG hati tahunan"],
        ["**Diklorometana**", "Anestesi di ruang tertutup, CO-Hb, infark", "DCM urin 0,3 mg/L; CO-Hb 3,5% (akhir shift)", "**O2 100%/HBO**, lebih lama (t1/2 ± 13 jam)"],
        ["**CCl4**", "Nekrosis hati sentrilobular lalu AKI", "Tidak ada BEI", "NAC (bukti terbatas); jangan epinefrin dan alkohol"],
        ["**TCE**", "Kanker ginjal, trigeminal, DRESS, degreaser's flush", "TCA urin 15 mg/L; TCOH 0,5 mg/L (akhir shift akhir pekan)", "Hentikan pajanan; steroid (hipersensitivitas)"],
        ["**PCE**", "Diskromatopsia, ensefalopati", "PCE napas 3 ppm / darah (sebelum shift berikutnya)", "Suportif"],
        ["**CS2**", "Psikosis, polineuropati, PJK, mikroaneurisma retina", "TTCA urin 0,5 mg/g (B, Ns)", "Benzodiazepin; **hindari haloperidol**"],
        ["**Aseton / DIPE / etil asetat**", "Iritasi, SSP ringan; DIPE: peroksida", "Aseton urin 25 mg/L (B, Ns)", "Suportif"],
    ]
    M.add_table(doc, rows, font=8, widths=[2.4, 4.9, 5.3, 4.2])


def penutup_3b(doc):
    doc.add_heading("6.4 Lafaz OSCE: satu kalimat per zat (3B)", level=2)
    for t in [
        "**Benzena:** \"Pada pajanan rendah pilih S-PMA karena t,t-MA terganggu asam sorbat; benzena tanpa ambang, tanpa antidot, dan NAB Indonesia 0,5 ppm jauh di atas TLV ACGIH 0,02 ppm.\"",
        "**Toluena:** \"Toluena adalah ototoksikan yang sinergis dengan bising; pantau o-kresol urin, bukan asam hipurat, dan toluena darah diambil sebelum shift terakhir pekan.\"",
        "**n-Heksana:** \"Neuropati dengan foot drop pada pekerja sepatu; 2,5-HD urin sering sudah negatif sehingga diagnosis dari riwayat dan NCV/EMG, dan pekerja bisa memburuk 1 sampai 4 bulan (coasting).\"",
        "**MEK:** \"MEK sendiri tidak neurotoksik tetapi menginduksi CYP2E1 sehingga memperkuat 2,5-HD dari n-heksana atau MBK.\"",
        "**Anilin/nitrobenzena:** \"Sianosis refrakter oksigen dengan SpO2 tetap 85% adalah methemoglobinemia: biru metilen 1 sampai 2 mg/kg setelah memastikan G6PD, dan transfusi tukar bila defisiensi G6PD.\"",
        "**Metanol:** \"Penglihatan badai salju dengan Kussmaul setelah miras oplosan: fomepizol, bikarbonat, asam folinat, dan hemodialisis; asam format, bukan formaldehida, penyebab buta.\"",
        "**Etilen glikol:** \"Mabuk tanpa bau alkohol lalu asidosis anion gap dan AKI dengan kristal oksalat: fomepizol, tiamin dan piridoksin, dan kalsium hanya bila hipokalsemia simtomatik.\"",
        "**Formaldehida:** \"Teknisi patologi dengan asma sensitisasi dinyatakan tidak laik permanen; tidak ada BEI, ceiling Indonesia 0,3 ppm, dan karsinogen site-of-contact (nasofaring).\"",
        "**VCM:** \"Trombosit rendah dengan splenomegali pada pekerja PVC adalah hipertensi portal non-sirotik; surveilans USG hati dan ingat angiosarkoma hati.\"",
        "**Diklorometana:** \"CO-Hb dari DCM terus naik beberapa jam setelah evakuasi dengan t1/2 ± 13 jam sehingga oksigen diberikan lebih lama dan respirator kartrid tidak boleh dipakai.\"",
        "**TCE:** \"Mati rasa wajah (trigeminal), kanker ginjal, dan flushing setelah alkohol adalah TCE; TCA urin diambil Jumat karena t1/2 70 sampai 100 jam.\"",
        "**PCE:** \"Pekerja dry cleaning dideteksi dini dengan D-15 (biru-kuning); PCE darah atau napas diambil sebelum shift berikutnya karena depot lemak.\"",
        "**CS2:** \"Mikroaneurisma retina pada pekerja viskosa muda adalah vaskulopati CS2 yang menandai risiko koroner; hindari haloperidol karena dopamin sudah tinggi.\"",
    ]:
        bullets(doc, [t])

    doc.add_heading("6.5 Vignette MCQ (3B)", level=2)
    para(doc, "Skenario fiktif untuk latihan.", size=9.5, italic=True)
    V = [
        ("Vignette 5 (Anilin)",
         "Pekerja pabrik MDI 33 tahun terkena percikan anilin di lengan. Satu jam kemudian asimtomatik, lalu 4 jam kemudian sianosis, SpO2 85% tidak naik dengan O2 100%, PaO2 normal, darah cokelat. Langkah paling tepat setelah dekontaminasi?",
         ["A. Antibiotik", "B. Co-oximetry, lalu biru metilen 1 sampai 2 mg/kg IV bila G6PD tidak defisien", "C. Atropin", "D. Hidroksokobalamin", "E. Intubasi dan observasi saja"],
         "B. Methemoglobinemia (saturation gap); biru metilen dengan batas kumulatif 7 mg/kg; pada defisiensi G6PD gunakan transfusi tukar."),
        ("Vignette 6 (Metanol)",
         "Pria 28 tahun, 18 jam setelah miras oplosan, pandangan seperti badai salju, napas Kussmaul, pH 7,18, anion gap tinggi, osmolal gap sedang. Terapi yang paling tepat?",
         ["A. Observasi 24 jam", "B. Fomepizol, bikarbonat, asam folinat, dan persiapan hemodialisis", "C. Tiamin dan piridoksin", "D. Kalsium glukonat", "E. N-asetilsistein"],
         "B. Metanol: blokade ADH, koreksi asidosis, folat untuk eliminasi format; dialisis untuk gangguan visus dan pH rendah. Tiamin dan B6 untuk etilen glikol."),
        ("Vignette 7 (Diklorometana)",
         "Pengecat bak mandi 45 tahun, perokok, pingsan di kamar mandi tertutup. CO-Hb awal 12% dan stabil dengan O2 100%; IGD berencana memulangkan 6 jam kemudian. Penilaian paling tepat?",
         ["A. Pulang, CO-Hb sudah turun", "B. Pantau CO-Hb serial minimal 24 jam karena DCM terus menjadi CO dan CO-Hb dapat naik lagi", "C. Berikan epinefrin", "D. Kelasi", "E. Rujuk psikiatri"],
         "B. Depot lemak melepas DCM; t1/2 CO-Hb ± 13 jam (6 jam dengan O2 100%); oksigen diberikan lebih lama."),
    ]
    for title, stem, opts, ans in V:
        para(doc, "**" + title + "**", after=2)
        para(doc, stem, after=2)
        for o in opts:
            para(doc, o, after=0).paragraph_format.left_indent = Cm(0.8)
        para(doc, "**Jawaban:** " + ans, size=9.5, after=8)
