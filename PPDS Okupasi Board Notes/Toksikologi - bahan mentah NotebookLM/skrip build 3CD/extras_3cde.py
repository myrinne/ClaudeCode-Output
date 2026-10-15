from docx.shared import Pt, Cm
import mdconv as M
from extras import para, bullets, callout


def intro_3c(doc):
    para(doc, "Struktur sama dengan 3A dan 3B: tabel 13 poin per zat, pearl dan jebakan ujian (poin 14), dan **satu infografik per zat** (sumber, karakteristik, toksikokinetik dan mekanisme, efek akut dan kronis, biomarker dan nilai batas, tatalaksana dan clinical pearl). Sumber: notebook NotebookLM \"Toxicology - Pesticide\" (39 sumber). Poin 12 tidak merinci urutan langkah diagnosis PAK. Nilai yang ditandai \"cek ulang\" ada di bagian 7.", size=10)
    M.add_table(doc, [
        ["Kelompok (outline 3C)", "Zat pada catatan ini"],
        ["Organofosfat", "Diklorvos, paration, malation, klorpirifos (kolinesterase eritrosit vs plasma, sindrom intermediate, OPIDN)"],
        ["Organoklorin", "DDT, lindane/HCH, siklodiena (aldrin, dieldrin, endosulfan, klordan)"],
        ["Karbamat", "Karbaril, karbofuran, metomil, propoksur (inhibisi reversibel)"],
        ["Piretroid", "Permetrin, sipermetrin, deltametrin, sihalotrin (parestesia)"],
        ["Paraquat", "Paraquat dengan diquat sebagai pembanding (fibrosis paru, tes ditionit urin)"],
        ["Rodentisida / fumigan", "Fosfin dan aluminium fosfida; metil bromida"],
    ], font=9, widths=[4.5, 12.3])
    callout(doc,
            "**Jebakan lintas pestisida.** (1) **Kolinesterase hanya untuk OP dan karbamat;** organoklorin, piretroid, paraquat, fosfin, dan metil bromida punya ChE normal. "
            "(2) **Atropin dan oksim hanya untuk OP/karbamat:** atropin pada keracunan organoklorin dan piretroid memperburuk; oksim pada karbaril relatif dihindari. "
            "(3) **Suksinilkolin dilarang pada OP dan karbamat.** (4) **Epinefrin dilarang pada organoklorin** (sensitisasi miokard). "
            "(5) **Oksigen tinggi dilarang pada paraquat;** lavage air biasa dilarang pada aluminium fosfida. (6) **Metil bromida: periode laten 2 sampai 24 jam, rawat 24 sampai 48 jam.**",
            fill="FDE9D9")


def quick_table_3c(doc):
    doc.add_heading("5C. QUICK TABLE AKHIR (3C. PESTISIDA)", level=1)
    para(doc, "Kolom: zat | gejala kunci | biomarker (sampel + waktu) | antidotum / terapi kunci.", size=9.5, italic=True)
    rows = [
        ["Zat", "Gejala kunci", "Biomarker (sampel + waktu)", "Antidotum / terapi"],
        ["**Organofosfat**", "Krisis kolinergik (bronkorea, bronkospasme, bradikardia), fasikulasi, kejang; IMS hari 1 sampai 4; OPIDN minggu 1 sampai 3", "AChE eritrosit ≥ 70% baseline individu (IPB/BEI, kapan saja); BuChE plasma; DAP urin; p-nitrofenol 0,5 mg/g (paration)", "O2 dahulu; **atropin 1 sampai 2 mg digandakan sampai paru kering**; **pralidoksim 30 mg/kg lalu 8 sampai 10 mg/kg/jam**; diazepam; **tanpa suksinilkolin**"],
        ["**Organoklorin**", "DDT: parestesia wajah + tremor kasar; siklodiena/lindane: kejang mendadak; kronis: endokrin, NHL, Parkinson", "Serum DDT/DDE per lemak; lindane darah > 20 µg/L; **tidak ada IPB/BEI, ChE tidak berguna**", "Diazepam, fenobarbital; **kolestiramin 4 g/6 jam**, arang aktif; **tanpa epinefrin, atropin, oksim, susu/minyak**"],
        ["**Karbamat**", "Krisis kolinergik singkat (6 sampai 24 jam); tanpa OPIDN, IMS jarang", "AChE eritrosit 70% baseline; **normal palsu (dekarbamilasi in vitro): sampel dingin 0 sampai 4°C**; 1-naftol urin (karbaril)", "Atropin; **oksim tidak rutin, hindari pada karbaril murni**, beri bila berat atau campuran OP; tanpa suksinilkolin"],
        ["**Piretroid**", "**Parestesia kulit wajah tanpa eritema**; tipe I tremor (sindrom T), tipe II salivasi, choreoathetosis, kejang (sindrom CS)", "**ChE normal;** 3-PBA, cis/trans-DCCA, DBCA urin akhir shift; tidak ada IPB", "**Vitamin E topikal;** diazepam; **tanpa atropin rutin**, tanpa oksim; hindari air dingin dan pasta seng oksida"],
        ["**Paraquat**", "Paraquat tongue, ATN, hepatotoksik, **fibrosis paru 7 sampai 21 hari**; kuku, epistaksis", "**Ditionit urin biru** (diquat hijau); plasma Proudfoot (jam 4 < 2 mg/L); SIPP", "Tanah Fuller/bentonit/arang aktif; **hindari O2 tinggi (SpO2 85 sampai 88%)**; hemoperfusi hanya 2 sampai 4 jam; siklofosfamid + metilprednisolon; tanpa antidot"],
        ["**Fosfin / AlP**", "Syok refrakter, aritmia, EKG ST meniru STEMI, ARDS, asidosis laktat; napas bau ikan", "Uji perak nitrat (noda hitam); COHb palsu; tidak ada IPB", "**Tanpa antidot;** minyak kelapa + NaHCO3 (bukan air); bikarbonat, MgSO4, NAC, IABP/ECMO; **tanpa trombolitik**"],
        ["**Metil bromida**", "**Laten 2 sampai 24 jam**, lalu edema paru dan kejang; ataksia serebelar; bula raksasa", "Bromida serum (< 5 normal; > 12 sampai 15 pajanan; > 150 sampai 200 fatal); t½ 9 sampai 12 hari", "**Tanpa antidot; rawat 24 sampai 48 jam;** diazepam, NAC; NaCl 0,9% + furosemid (eliminasi bromida)"],
    ]
    M.add_table(doc, rows, font=8, widths=[2.3, 4.9, 5.2, 4.4])


def penutup_3c(doc):
    doc.add_heading("6.6 Lafaz OSCE: satu kalimat per zat (3C)", level=2)
    for t in [
        "**Organofosfat:** \"Setelah oksigenasi, atropin 1 sampai 2 mg digandakan tiap 5 menit sampai paru kering (bukan sampai pupil midriasis), pralidoksim 30 mg/kg lalu infus, diazepam untuk kejang, dan suksinilkolin dilarang; pekerja dipindah bila AChE eritrosit turun lebih dari 30% baseline dan kembali bila pulih 80%.\"",
        "**Organoklorin:** \"Bukan inhibitor kolinesterase, sehingga atropin dan oksim salah; kontrol kejang dengan benzodiazepin, hindari epinefrin dan susu/minyak, dan beri arang aktif serta kolestiramin.\"",
        "**Karbamat:** \"Kolinesterase bisa normal palsu karena dekarbamilasi in vitro sehingga sampel didinginkan dan diagnosis klinis; oksim tidak rutin dan dapat memperburuk karbaril.\"",
        "**Piretroid:** \"Parestesia wajah tanpa eritema dengan kolinesterase normal; oleskan vitamin E, jangan atropin walau berliur, dan surveilans memakai 3-PBA urin.\"",
        "**Paraquat:** \"Hindari oksigen kecuali hipoksemia mengancam jiwa (target SpO2 85 sampai 88%), beri tanah Fuller atau arang aktif, hemoperfusi hanya dalam 2 sampai 4 jam, dan tes ditionit urin biru; diquat tidak merusak paru.\"",
        "**Fosfin/AlP:** \"Jangan lavage dengan air; beri minyak kelapa dan bikarbonat, MgSO4, NAC, dukungan IABP/ECMO, tanpa trombolitik karena EKG elevasi ST adalah miokarditis; tenaga IGD memakai respirator gas.\"",
        "**Metil bromida:** \"Pasien yang tampak sehat setelah pajanan tetap dirawat 24 sampai 48 jam karena laten 2 sampai 24 jam; tanpa antidot, SCBA wajib, dan sarung tangan karet memperparah bula.\"",
    ]:
        bullets(doc, [t])

    doc.add_heading("6.7 Vignette MCQ (3C)", level=2)
    para(doc, "Skenario fiktif untuk latihan.", size=9.5, italic=True)
    V = [
        ("Vignette 8 (Organofosfat)",
         "Petani 40 tahun, 2 jam setelah menyemprot klorpirifos: berkeringat, miosis, mengi, nadi 48/menit, fasikulasi. Setelah atropin dan oksim ia stabil, tetapi hari ke-3 tidak mampu mengangkat kepala, ptosis, dan napas dangkal dengan sensorik normal. Diagnosis dan tindakan?",
         ["A. OPIDN; berikan atropin lebih banyak", "B. Sindrom intermediate; dukungan ventilasi, atropin tidak membantu", "C. Krisis miastenik; neostigmin", "D. Sindrom Guillain-Barre; imunoglobulin", "E. Toksisitas atropin"],
         "B. IMS muncul 24 sampai 96 jam setelah krisis kolinergik (kelemahan proksimal, fleksor leher, kranial, napas; sensorik utuh) dan butuh ventilator; OPIDN baru minggu 1 sampai 3."),
        ("Vignette 9 (Piretroid)",
         "Penyemprot kelambu 28 tahun mengeluh wajah terbakar dan berliur 2 jam setelah bekerja. HR 112, pupil normal, paru bersih, kolinesterase 100%. Terapi paling tepat?",
         ["A. Atropin 2 mg IV berulang", "B. Pralidoksim", "C. Oleskan vitamin E topikal, hindari air dingin, istirahat dari semprot", "D. Cuci wajah dengan air dingin berulang", "E. Pasta seng oksida"],
         "C. Piretroid (parestesia kutaneus, salivasi tipe II); bukan kolinergik sehingga atropin dan oksim tidak ada indikasi."),
        ("Vignette 10 (Paraquat)",
         "Pria 35 tahun menelan satu teguk herbisida 3 jam lalu; luka orofaring, SpO2 94% udara ruang. Urin + natrium ditionit menjadi biru tua. Langkah paling tepat?",
         ["A. Oksigen 100% non-rebreather", "B. Tanah Fuller atau arang aktif segera, hindari O2 tinggi, pertimbangkan hemoperfusi dini, dan nilai prognosis", "C. Atropin dan pralidoksim", "D. Lavage air bervolume besar", "E. Rawat jalan"],
         "B. Paraquat (biru pada ditionit): adsorben, hindari oksigen berlebihan, hemoperfusi hanya dalam 2 sampai 4 jam; transplantasi paru gagal."),
        ("Vignette 11 (Metil bromida)",
         "Teknisi fumigasi kontainer terkena kebocoran selang metil bromida 4 jam lalu. Kini tanda vital normal, paru bersih, meminta pulang. Tindakan paling tepat?",
         ["A. Pulang dan kontrol 1 minggu", "B. Rawat dan observasi 24 sampai 48 jam karena periode laten", "C. Berikan dimerkaprol", "D. Atropin", "E. Pulang dengan steroid oral"],
         "B. Laten 2 sampai 24 jam (sampai 48 jam) sebelum edema paru dan kejang; tanpa antidot, suportif, bromida serum, NAC bila perlu."),
    ]
    for title, stem, opts, ans in V:
        para(doc, "**" + title + "**", after=2)
        para(doc, stem, after=2)
        for o in opts:
            para(doc, o, after=0).paragraph_format.left_indent = Cm(0.8)
        para(doc, "**Jawaban:** " + ans, size=9.5, after=8)


def quality_extra_3c(doc):
    bullets(doc, [
        "**3C, NAB fosfin (Permenaker 5/2018):** NotebookLM menyebut NAB 0,3 ppm dan KTD 1 ppm, sedangkan ACGIH TLV-TWA 0,05 ppm dan STEL 0,15 ppm; cek ulang Lampiran 3 Permenaker.",
        "**3C, nomor butir Perpres 7/2019 dari NotebookLM:** pestisida butir 35, fosfor butir 3, halogen hidrokarbon butir 24 (organoklorin) atau butir 11 (metil bromida), napas I.c.2; tidak konsisten antar jawaban dan belum diverifikasi ke lampiran.",
        "**3C, klasifikasi IARC:** aldrin dan dieldrin tertulis Grup 2B oleh NotebookLM (pustaka lain: Grup 2A); lindane Grup 1 dan DDT 2A sesuai IARC 2015 (cek ulang).",
        "**3C, bukti terbatas:** vitamin E topikal untuk parestesia piretroid (angka 98% dari studi lama), imunosupresi siklofosfamid-metilprednisolon pada paraquat (bukti beragam), minyak kelapa untuk keracunan AlP, dan target SpO2 85 sampai 88% pada paraquat.",
        "**3C, tidak ada infografik tambahan di luar tujuh zat;** angka pada infografik dihasilkan NotebookLM dari prompt dan perlu dicek terhadap tabel.",
    ], size=10)
