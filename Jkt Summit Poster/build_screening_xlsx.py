"""Build the title/abstract screening workbooks from records.json + decisions_r1.py.

Outputs:
  Screening_TA_master.xlsx  - R1 (Claude) decisions + R2 + consensus columns, PRISMA counts, kappa
  Screening_TA_blind_R2.xlsx - same records and criteria, no R1 columns (give to Reviewer 2)
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from decisions_r1 import D

HERE = Path(__file__).parent
data = json.loads((HERE / "records.json").read_text(encoding="utf-8"))
R, DUPS, COUNTS = data["records"], data["duplicates"], data["counts"]

HEAD = PatternFill("solid", fgColor="1F4E78")
HFONT = Font(bold=True, color="FFFFFF")
R1FILL = PatternFill("solid", fgColor="DDEBF7")
R2FILL = PatternFill("solid", fgColor="FFF2CC")
CFILL = PatternFill("solid", fgColor="E2EFDA")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
WRAP = Alignment(wrap_text=True, vertical="top")

REASONS = [
    ("E1", "Population/setting: not workers undergoing periodic, pre-placement/pre-employment or statutory occupational health examination (incl. general-population or patient health check-ups, HR/personnel selection)"),
    ("E2", "Concept: no AI/ML/DL/NLP/LLM or computerised CDS (conventional statistics only)"),
    ("E3", "AI not applied to examination data (exposure measurements, wearables/sensors, surveys, sickness absence, claims, injury narratives, social media)"),
    ("E4", "Prognostic only: predicts FUTURE disease onset/progression; no classification of the current examination result"),
    ("E5", "Not primary research: review, editorial, commentary, perspective, guidance, protocol, chapter, conference panel"),
    ("E6", "Non-human (animal, plant, veterinary)"),
    ("E7", "No performance against a reference standard (unsupervised clustering/association/feature importance only)"),
]
CRITERIA = [
    ("Review", "AI-based clinical decision support in periodic occupational health examinations: diagnostic performance and implementation barriers (JBI scoping review, PRISMA-ScR)"),
    ("P - Population", "Working-age people undergoing periodic or pre-placement/pre-employment occupational health examinations (any industry, any country)"),
    ("C - Concept", "AI-based CDS (ML classifiers, deep learning incl. imaging, NLP/LLM) applied to examination modalities (chest radiography/CT, spirometry, audiometry, laboratory, ECG, multimodal) for result classification, abnormality detection or examination-level categorisation (incl. fitness-for-work)"),
    ("C - Context", "Occupational health screening settings; comparator = physician/specialist reference standard or clinically validated diagnosis"),
    ("Outcomes", "Primary: Se, Sp, PPV, NPV, accuracy, AUC, Cohen's kappa. Secondary: implementation barriers (HOT framework)"),
    ("Stage rule", "Title/abstract stage is INCLUSIVE: if P and C plausibly met but outcome/reference unclear -> M (Maybe), decided at full text"),
    ("Decision codes", "I = include to full text | M = maybe, retrieve full text | E = exclude (give ONE reason code, first that applies in order E5 > E6 > E1 > E2 > E3 > E4 > E7)"),
    ("Not restricted", "No date or language limit was applied at this stage (one Japanese-language record, #92)"),
    ("Reviewer 2", "Screen independently in Screening_TA_blind_R2.xlsx without looking at R1. Then paste your Decision/Reason/Note columns into the master's R2 columns (same row order, check No.)"),
    ("Conflicts", "Master column 'Agree?' flags disagreements on I/M vs E. Resolve by discussion; record final call in Consensus columns. PRISMA counts and kappa update automatically"),
    ("Searches", "PubMed (109), Scopus (190, from emailed document list because RIS export failed), ProQuest (32), IEEE Xplore (36). Abstracts missing from PubMed/IEEE exports were fetched from NCBI E-utilities, OpenAlex and Semantic Scholar"),
]


def header(ws, cols, fills=None):
    for i, (name, width) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.fill = (fills or {}).get(name, HEAD)
        c.font = HFONT if c.fill is HEAD else Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "H2"
    ws.row_dimensions[1].height = 32


def criteria_sheet(wb):
    ws = wb.create_sheet("Criteria")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 120
    r = 1
    ws.cell(r, 1, "Eligibility criteria - title/abstract screening").font = Font(bold=True, size=13)
    r += 2
    for k, v in CRITERIA:
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v).alignment = WRAP
        r += 1
    r += 1
    ws.cell(r, 1, "Exclusion reason codes").font = Font(bold=True, size=12)
    r += 1
    for k, v in REASONS:
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v).alignment = WRAP
        r += 1
    return ws


def screening_sheet(wb, blind):
    ws = wb.active
    ws.title = "Screening TA"
    base = [("No.", 6), ("Source", 11), ("Also in", 12), ("Record ID", 16), ("Year", 6), ("First authors", 18),
            ("Journal / venue", 22), ("Title", 55), ("Abstract", 90)]
    r1 = [("R1 Decision", 9), ("R1 Reason", 8), ("R1 Note", 40)]
    r2 = [("R2 Decision", 9), ("R2 Reason", 8), ("R2 Note", 30)]
    cons = [("Agree?", 8), ("Consensus Decision", 11), ("Consensus Reason", 10), ("Consensus Note", 30)]
    cols = base + ([] if blind else r1) + r2 + ([] if blind else cons)
    fills = {n: R1FILL for n, _ in r1} | {n: R2FILL for n, _ in r2} | {n: CFILL for n, _ in cons}
    header(ws, cols, fills)
    idx = {name: i + 1 for i, (name, _) in enumerate(cols)}
    for row, rec in enumerate(R, 2):
        vals = [rec["no"], rec["source"], rec.get("also_in", ""), rec["id"], rec["year"], rec["authors"],
                rec["journal"], rec["title"], rec["abstract"] or "[no abstract available - screen on title]"]
        if not blind:
            dec, rsn, note = D[rec["no"]]
            vals += [dec, rsn, note]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row, col, v)
            c.alignment = WRAP
            c.border = THIN
        ws.row_dimensions[row].height = 90
        if not blind:
            L = lambda n: f"{get_column_letter(idx[n])}{row}"
            ws.cell(row, idx["Agree?"], f'=IF({L("R2 Decision")}="","",IF(({L("R1 Decision")}="E")=({L("R2 Decision")}="E"),"Yes","CONFLICT"))')
    last = len(R) + 1
    dv_dec = DataValidation(type="list", formula1='"I,M,E"', allow_blank=True)
    dv_rsn = DataValidation(type="list", formula1='"' + ",".join(k for k, _ in REASONS) + '"', allow_blank=True)
    ws.add_data_validation(dv_dec)
    ws.add_data_validation(dv_rsn)
    targets = [("R2 Decision", "R2 Reason")] + ([] if blind else [("Consensus Decision", "Consensus Reason")])
    for d_col, r_col in targets:
        dv_dec.add(f"{get_column_letter(idx[d_col])}2:{get_column_letter(idx[d_col])}{last}")
        dv_rsn.add(f"{get_column_letter(idx[r_col])}2:{get_column_letter(idx[r_col])}{last}")
    if not blind:
        a = get_column_letter(idx["Agree?"])
        ws.conditional_formatting.add(f"{a}2:{a}{last}", FormulaRule(formula=[f'{a}2="CONFLICT"'], fill=PatternFill("solid", fgColor="F8CBAD")))
        d = get_column_letter(idx["R1 Decision"])
        ws.conditional_formatting.add(f"{d}2:{d}{last}", FormulaRule(formula=[f'{d}2="I"'], fill=PatternFill("solid", fgColor="A9D08E")))
        ws.conditional_formatting.add(f"{d}2:{d}{last}", FormulaRule(formula=[f'{d}2="M"'], fill=PatternFill("solid", fgColor="FFE699")))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
    return ws, idx, last


def dup_sheet(wb):
    ws = wb.create_sheet("Duplicates removed")
    cols = [("Source", 11), ("Record ID", 18), ("Title", 80), ("Duplicate of (kept No./ID)", 22)]
    header(ws, cols)
    ws.freeze_panes = "A2"
    kept = {r["id"]: r["no"] for r in R}
    for row, d in enumerate(DUPS, 2):
        for col, v in enumerate([d["source"], d["id"], d["title"], f'#{kept.get(d["duplicate_of"], "?")} ({d["duplicate_of"]})'], 1):
            ws.cell(row, col, v).alignment = WRAP


def prisma_sheet(wb, idx, last):
    ws = wb.create_sheet("PRISMA counts")
    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 60
    S = "'Screening TA'!"
    col = lambda n: f"{S}${get_column_letter(idx[n])}$2:${get_column_letter(idx[n])}${last}"
    cd, cr, r1, r2 = col("Consensus Decision"), col("Consensus Reason"), col("R1 Decision"), col("R2 Decision")
    rows = [
        ("IDENTIFICATION", None, None),
        ("Records from PubMed/MEDLINE", COUNTS["PubMed"], ""),
        ("Records from Scopus", COUNTS["Scopus"], "From emailed Scopus document list"),
        ("Records from ProQuest", COUNTS["ProQuest"], ""),
        ("Records from IEEE Xplore", COUNTS["IEEE Xplore"], ""),
        ("Total records identified", "=SUM(B3:B6)", ""),
        ("Duplicates removed", len(DUPS), "DOI or normalised-title match"),
        ("SCREENING (title/abstract)", None, None),
        ("Records screened", f"=COUNTA({S}$A$2:$A${last})", ""),
        ("Records excluded (consensus)", f'=COUNTIF({cd},"E")', "Fills once Consensus Decision is complete"),
    ]
    for code, desc in REASONS:
        rows.append((f"   {code}: {desc[:70]}", f'=COUNTIFS({cd},"E",{cr},"{code}")', ""))
    rows += [
        ("Reports sought for retrieval (consensus I + M)", f'=COUNTIF({cd},"I")+COUNTIF({cd},"M")', "Next stage: full-text screening"),
        ("Consensus decisions still blank", f'=B10-COUNTA({cd})', "Should reach 0"),
        ("", None, None),
        ("INTER-RATER AGREEMENT (I/M vs E, binary)", None, None),
        ("Both screened", f'=SUMPRODUCT(({r1}<>"")*({r2}<>""))', ""),
        ("Both keep (I/M)", f'=SUMPRODUCT((({r1}="I")+({r1}="M"))*(({r2}="I")+({r2}="M")))', ""),
        ("Both exclude (E)", f'=SUMPRODUCT(({r1}="E")*({r2}="E"))', ""),
        ("R1 keep, R2 exclude", f'=SUMPRODUCT((({r1}="I")+({r1}="M"))*({r2}="E"))', ""),
        ("R1 exclude, R2 keep", f'=SUMPRODUCT(({r1}="E")*(({r2}="I")+({r2}="M")))', ""),
    ]
    start = len(rows) + 1
    po = f"(B{start-3}+B{start-2})/B{start-4}"
    pe = f"((B{start-3}+B{start-1})*(B{start-3}+B{start})+(B{start-2}+B{start})*(B{start-2}+B{start-1}))/B{start-4}^2"
    rows += [
        ("Observed agreement (Po)", f"=IFERROR({po},\"\")", ""),
        ("Cohen's kappa", f"=IFERROR(({po}-{pe})/(1-{pe}),\"\")", "Report in Methods: kappa for title/abstract screening"),
        ("", None, None),
        ("REVIEWER 1 (Claude) PRELIMINARY TALLY", None, None),
        ("R1 include (I)", f'=COUNTIF({r1},"I")', ""),
        ("R1 maybe (M)", f'=COUNTIF({r1},"M")', ""),
        ("R1 exclude (E)", f'=COUNTIF({r1},"E")', ""),
    ]
    for i, (a, b, c) in enumerate(rows, 1):
        ws.cell(i + 1, 1, a)
        if b is None and a:
            ws.cell(i + 1, 1).font = Font(bold=True, color="1F4E78")
        if b is not None:
            ws.cell(i + 1, 2, b)
        if c:
            ws.cell(i + 1, 3, c).alignment = WRAP


def main():
    for blind, name in [(False, "Screening_TA_master.xlsx"), (True, "Screening_TA_blind_R2.xlsx")]:
        wb = Workbook()
        _, idx, last = screening_sheet(wb, blind)
        criteria_sheet(wb)
        if not blind:
            dup_sheet(wb)
            prisma_sheet(wb, idx, last)
        wb.save(HERE / name)
        print("saved", name)


if __name__ == "__main__":
    main()
