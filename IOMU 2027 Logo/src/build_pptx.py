"""Build the IOMU 2027 logo voting deck (palette B, 3 concepts)."""
import os
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt, Emu

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
IMG = os.path.join(ROOT, "ai-concepts")
OUT = os.path.join(ROOT, "IOMU 2027 Logo Concepts v3.pptx")

EMERALD = RGBColor(0x21, 0x49, 0x40)
GOLD = RGBColor(0xC6, 0x9C, 0x54)
CREAM = RGBColor(0xF4, 0xE9, 0xD7)
INK = RGBColor(0x2B, 0x2B, 0x2B)
HEAD, BODY = "Georgia", "Calibri"

CONCEPTS = [
    ("1", "Healthy Worker", os.path.join("Revised", "IOMU 2027 Emblem and Wordmark.png"), [
        ("Figure", "An abstract worker with arms raised: a healthy, thriving workforce"),
        ("Hard-hat head", "Occupational safety and the worker as the centre of the specialty"),
        ("Rod of Asclepius", "Medicine, forming the figure's backbone"),
        ("Batik ribbon", "Solo batik in flowing emerald and gold bands"),
        ("Ornament in the O", "Gold crest ornament inspired by Javanese carving"),
        ("Serif wordmark", "Classic, formal, matches a scientific meeting"),
    ]),
    ("2", "Kawung Cross", "iomu2027_chatgpt_c2_kawung-cross_B.png", [
        ("Four kawung petals", "Arranged as a medical cross: health built on Javanese heritage"),
        ("Batik engraving", "Parang pattern inside each petal"),
        ("Hard hat + rod", "Occupational safety and medicine at the centre"),
        ("Crest above the M", "Ornament inspired by the carved crest of Keraton Surakarta"),
        ("Sans-serif wordmark", "Modern, strong, highly legible at small sizes"),
    ]),
    ("3", "Gunungan", "iomu2027_merge_gunungan_serif-type_B_v1.png", [
        ("Gunungan", "Wayang kulit tree of life; Solo is a centre of wayang tradition"),
        ("Hard hat + rod", "Occupational safety and medicine inside the gunungan"),
        ("Batik bands", "Parang engraving with gold scrollwork"),
        ("Kawung in the O", "Solo batik; order and integrity"),
        ("Serif wordmark", "Classic, formal, matches a scientific meeting"),
    ]),
]


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(slide, x, y, w, h, s, size, color, font=BODY, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = s
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.name = font
    r.font.bold = bold
    return tb


def rule(slide, x, y, w, color=GOLD, h=Pt(2)):
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    ln.fill.solid()
    ln.fill.fore_color.rgb = color
    ln.line.fill.background()
    return ln


def picture_fit(slide, path, x, y, w, h):
    """Place an image inside the box (x, y, w, h), keeping aspect, centred."""
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih)
    pw, ph = int(iw * s), int(ih * s)
    return slide.shapes.add_picture(path, Emu(x + (w - pw) // 2), Emu(y + (h - ph) // 2), width=Emu(pw), height=Emu(ph))


def footer(slide, s="Concept draft, AI-generated. The chosen logo will be redrawn as a clean vector file."):
    text(slide, Inches(0.6), Inches(7.0), Inches(12.1), Inches(0.35), s, 10, RGBColor(0x8A, 0x80, 0x70))


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    # title
    s = prs.slides.add_slide(blank)
    bg(s, EMERALD)
    text(s, Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.2), "IOMU 2027 Logo Concepts", 48, CREAM, HEAD)
    rule(s, Inches(0.95), Inches(3.55), Inches(2.2))
    text(s, Inches(0.9), Inches(3.8), Inches(11.5), Inches(0.6),
         "19th Indonesian Occupational Medicine Updates · Solo · PERDOKI", 20, GOLD, BODY)
    text(s, Inches(0.9), Inches(4.5), Inches(11.5), Inches(0.6),
         "Three candidates for voting · Palette: Emerald Green & Gold", 16, CREAM, BODY)

    # one slide per concept
    for num, name, img, notes in CONCEPTS:
        s = prs.slides.add_slide(blank)
        bg(s, RGBColor(0xFF, 0xFF, 0xFF))
        text(s, Inches(0.6), Inches(0.35), Inches(9), Inches(0.7), f"Concept {num} · {name}", 30, EMERALD, HEAD)
        rule(s, Inches(0.62), Inches(1.05), Inches(1.6))
        picture_fit(s, os.path.join(IMG, img), Inches(0.6), Inches(1.4), Inches(8.8), Inches(5.3))
        panel = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.6), Inches(1.75), Inches(3.15), Inches(4.6))
        panel.fill.solid()
        panel.fill.fore_color.rgb = CREAM
        panel.line.fill.background()
        tb = s.shapes.add_textbox(Inches(9.75), Inches(1.95), Inches(2.85), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, (k, v) in enumerate(notes):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(6)
            for t, bold, col in ((k + ": ", True, EMERALD), (v, False, INK)):
                r = p.add_run()
                r.text = t
                r.font.bold = bold
                r.font.size = Pt(12)
                r.font.color.rgb = col
                r.font.name = BODY
        footer(s)

    # voting overview
    s = prs.slides.add_slide(blank)
    bg(s, RGBColor(0xFF, 0xFF, 0xFF))
    text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.7), "Vote: which logo for IOMU 2027?", 30, EMERALD, HEAD)
    rule(s, Inches(0.62), Inches(1.05), Inches(1.6))
    w, h = Inches(4.0), Inches(3.0)
    for i, (num, name, img, _) in enumerate(CONCEPTS):
        x = Inches(0.6) + i * (w + Inches(0.17))
        picture_fit(s, os.path.join(IMG, img), x, Inches(1.9), w, h)
        text(s, x, Inches(5.05), w, Inches(0.5), f"{num} · {name}", 18, EMERALD, HEAD, align=PP_ALIGN.CENTER)
    footer(s)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
